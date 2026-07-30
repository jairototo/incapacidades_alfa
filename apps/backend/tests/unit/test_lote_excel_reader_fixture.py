# tests/unit/test_lote_excel_reader_fixture.py
"""Proves the password-protected xlsx test fixture is genuine.

This does NOT test application code -- the real decryption module
(`lote_excel_reader.py`) lands in a later task (3.1). This task only
guarantees future tasks have a working, real encrypted xlsx fixture to
build against, per Phase 0 Task 0.3.

The fixture (tests/fixtures/afp_formalizacion_encrypted.xlsx) was generated
by building a plain openpyxl workbook (sheet "FORMALIZACION", one header row
+ one dummy data row) and encrypting it with msoffcrypto's
`OfficeFile.encrypt()`. The fixture password below is for this test fixture
only -- it is NOT the real AFP password (that lives in `PREVISIONALES_AFP_PASSWORD`,
set via environment/.env, never committed).
"""
import io
from pathlib import Path

import msoffcrypto
import openpyxl
import pytest

FIXTURE_PATH = Path(__file__).parent.parent / "fixtures" / "afp_formalizacion_encrypted.xlsx"
FIXTURE_PASSWORD = "test1234"


def test_fixture_file_exists():
    assert FIXTURE_PATH.exists(), f"missing fixture: {FIXTURE_PATH}"
    assert FIXTURE_PATH.stat().st_size > 0


def test_fixture_is_genuinely_encrypted_plain_openpyxl_fails():
    """openpyxl.load_workbook must reject the raw encrypted bytes.

    This proves the fixture isn't a plain xlsx masquerading as encrypted.
    """
    with pytest.raises(Exception):
        openpyxl.load_workbook(FIXTURE_PATH)


def test_fixture_decrypts_and_reads_dummy_data():
    with open(FIXTURE_PATH, "rb") as f:
        office_file = msoffcrypto.OfficeFile(f)
        office_file.load_key(password=FIXTURE_PASSWORD)

        decrypted_buf = io.BytesIO()
        office_file.decrypt(decrypted_buf)

    decrypted_buf.seek(0)
    wb = openpyxl.load_workbook(decrypted_buf)

    assert "FORMALIZACION" in wb.sheetnames
    ws = wb["FORMALIZACION"]

    header = [cell.value for cell in ws[1]]
    data_row = [cell.value for cell in ws[2]]

    assert header == ["TIPO_DOC", "NUMERO_DOC", "NOMBRE_AFILIADO", "MONTO"]
    assert data_row == ["CC", "123456789", "JUAN PEREZ TEST", 500000]


def test_fixture_wrong_password_raises():
    with open(FIXTURE_PATH, "rb") as f:
        office_file = msoffcrypto.OfficeFile(f)
        with pytest.raises(Exception):
            office_file.load_key(password="wrong-password", verify_password=True)
