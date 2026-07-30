"""Tests para el exportador ARPIS: agrupacion de salarios, radicado, tipo
de ingreso y construccion del libro Excel de cargue."""
import io
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app.services.previsionales.arpis_export import (
    DemasiadosGruposError,
    agrupar_salarios,
    construir_libro_arpis,
    mapear_tipo_ingreso,
    normalizar_radicado,
)
from app.services.previsionales.segmentacion import segmentar

# ---------------------------------------------------------------------------
# Fixtures de segmentos usadas por varios tests
# ---------------------------------------------------------------------------

s1 = segmentar(date(2026, 1, 1), date(2026, 1, 31))[0]
s2 = segmentar(date(2026, 2, 1), date(2026, 2, 28))[0]

# 4 segmentos consecutivos (uno por mes) para forzar 4 grupos de salario.
cuatro_segmentos = segmentar(date(2026, 1, 1), date(2026, 4, 30))
assert len(cuatro_segmentos) == 4


# ---------------------------------------------------------------------------
# agrupar_salarios
# ---------------------------------------------------------------------------


def test_salario_constante_es_un_solo_grupo():
    assert len(agrupar_salarios([(s, Decimal("1750905")) for s in segmentar(date(2026, 1, 1), date(2026, 7, 31))])) == 1


def test_dos_salarios_distintos_dos_grupos():
    g = agrupar_salarios([(s1, Decimal("100")), (s2, Decimal("200"))])
    assert [x.salario for x in g] == [Decimal("100"), Decimal("200")]


def test_mas_de_tres_grupos_alerta_no_pierde_datos():
    with pytest.raises(DemasiadosGruposError):  # la macro los metia callada en el grupo 3
        agrupar_salarios([(s, Decimal(i * 100)) for i, s in enumerate(cuatro_segmentos)])


def test_grupo_suma_dias_de_todos_sus_segmentos():
    segs = segmentar(date(2026, 1, 1), date(2026, 3, 31))  # 3 segmentos mensuales
    g = agrupar_salarios([(s, Decimal("500")) for s in segs])
    assert len(g) == 1
    assert g[0].dias == sum(s.dias for s in segs)


# ---------------------------------------------------------------------------
# normalizar_radicado
# ---------------------------------------------------------------------------


def test_radicado_15_digitos_recibe_cero():
    assert normalizar_radicado("103862022114100", None) == "0103862022114100"


def test_radicado_16_digitos_intacto():
    assert normalizar_radicado("4107413227182500", None) == "4107413227182500"


def test_radicado_no_numerico_sale_de_observacion():
    assert normalizar_radicado("N/A", "algo Rad No. 103862022114100 mas texto") == "0103862022114100"


def test_radicado_no_numerico_sin_patron_lanza():
    with pytest.raises(ValueError):
        normalizar_radicado("N/A", "una observacion sin el patron esperado")


def test_radicado_no_numerico_sin_observacion_lanza():
    with pytest.raises(ValueError):
        normalizar_radicado("N/A", None)


# ---------------------------------------------------------------------------
# mapear_tipo_ingreso
# ---------------------------------------------------------------------------


def test_tipo_ingreso_binario():
    assert mapear_tipo_ingreso("INICIAL") == (1, False, False)
    assert mapear_tipo_ingreso("TUTELA_I") == (1, False, True)
    assert mapear_tipo_ingreso("PRORROGA") == (2, False, False)
    assert mapear_tipo_ingreso("TUTELA_P") == (2, False, True)
    assert mapear_tipo_ingreso("AJUSTE") == (2, True, False)
    with pytest.raises(ValueError):
        mapear_tipo_ingreso("OTRO")  # "CASO NO DEFINIDO"


# ---------------------------------------------------------------------------
# construir_libro_arpis
# ---------------------------------------------------------------------------

fila = {
    "tipo_identificacion": "CC",
    "identificacion": "123456789",
    "dia_181": date(2026, 3, 1),
    "radicado": normalizar_radicado("103862022114100", None),
    "fecha_radicacion_afp": date(2026, 3, 5),
    "fecha_radicacion_alfa": date(2026, 3, 6),
    "tipo_ingreso": "INICIAL",
    "fecha_inicial": date(2026, 1, 1),
    "fecha_final": date(2026, 7, 31),
    "grupos_salario": agrupar_salarios(
        [(s, Decimal("1750905")) for s in segmentar(date(2026, 1, 1), date(2026, 7, 31))]
    ),
    "observacion": "Observacion AFP de prueba",
    "cie10_list": ["M545"],
    "observacion_causal": None,
}


def test_libro_tiene_encabezado_en_fila_3():
    ws = load_workbook(io.BytesIO(construir_libro_arpis([fila]))).active
    assert ws["A3"].value == "Tipo Identificación" and ws["A4"].value is not None


def test_libro_datos_desde_fila_4():
    ws = load_workbook(io.BytesIO(construir_libro_arpis([fila]))).active
    assert ws["B4"].value == "123456789"
    assert ws["D4"].value == "0103862022114100"
    assert ws["G4"].value == 1  # INICIAL -> codigo 1
    assert ws["J4"].value is None  # no es AJUSTE
    assert ws["K4"].value == Decimal("1750905")  # unico grupo de salario
    assert ws["L4"].value is None  # no hay segundo grupo
    assert ws["R4"].value == "M545"
    assert ws["S4"].value is None  # no es tutela


def test_libro_ajuste_y_tutela_marcan_columnas():
    fila_ajuste = dict(fila, tipo_ingreso="AJUSTE")
    ws = load_workbook(io.BytesIO(construir_libro_arpis([fila_ajuste]))).active
    assert ws["G4"].value == 2
    assert ws["J4"].value == 1

    fila_tutela = dict(fila, tipo_ingreso="TUTELA_I")
    ws = load_workbook(io.BytesIO(construir_libro_arpis([fila_tutela]))).active
    assert ws["S4"].value == 1


def test_libro_multiples_filas_avanzan_correctamente():
    ws = load_workbook(io.BytesIO(construir_libro_arpis([fila, fila]))).active
    assert ws["B4"].value == "123456789"
    assert ws["B5"].value == "123456789"
    assert ws["B6"].value is None
