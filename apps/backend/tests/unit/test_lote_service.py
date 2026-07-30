"""
Tests para `lote_service.py` (Task 3.1) -- el primer orquestador real del
módulo previsionales.

Usa `db_session` (Postgres async real, no mocks) y libros sintéticos en
memoria con la forma exacta de la hoja FORMALIZACION (columnas A-AA)
verificada contra el workbook real de producción. Los años usados en las
fechas de prueba (2026) deben tener SMLMV sembrado, porque los tests unit
usan `create_all` (no migraciones) y la tabla `smlmv_parametros` arranca
vacía.
"""
import io
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
import pytest_asyncio
from openpyxl import Workbook
from sqlalchemy import select

from app.core.exceptions import BadRequestException
from app.models.previsionales.incapacidad_previsional import IncapacidadPrevisional
from app.models.previsionales.periodo_previsional import PeriodoPrevisional
from app.models.previsionales.smlmv_parametros import SmlmvParametros
from app.services.previsionales.lote_service import lote_previsional_service

FIXTURE_PATH = Path(__file__).parent.parent / "fixtures" / "afp_formalizacion_encrypted.xlsx"

_ENCABEZADO_FORMALIZACION = [
    "Junio",
    "# IDENTIFICACION",
    "RADICADO",
    "FECHA RADICACION AFP",
    "TIPO INGRESO",
    "FECHA_INICIAL",
    "FECHA_FINAL",
    "No. DIAS",
    "DIAS ACUMULADOS",
    "DIAGNOSTICOS",
    "IBC",
    "SALARIO",
    "DIAS COTIZADOS",
    "FECHA DIA 181",
    "FECHA DIA 360",
    "VALOR",
    "FECHA RADICACION ASEGURADORA",
    "ASEGURADORA",
    "ESTADO AFILIADO",
    "DIA 181 ALFA",
    "FECHA CRIE ALFA",
    "OBSERVACIÓN ALFA",
    "OBSERVACIÓN PORVENIR",
    "VALOR PAGADO",
    "FECHA DE PAGO ASEGURADORA",
    "INSTANCIA JUDICIAL (cuando aplique)",
    "OBSERVACION JURIDICA (cuando aplique)",
]

# Filas válidas de referencia (orden de columnas A-AA):
_FILA_VALIDA_1 = [
    1, "1000000001", 123456789012345, date(2026, 6, 1), "INICIAL",
    "1/06/2026", "30/06/2026", 30, 30, "M545",
    1000000, 1000000, 30,
    date(2025, 1, 1), date(2025, 12, 1), 500000,
    date(2026, 6, 5), "Aseguradora Test", "VIGENTE",
    None, None, None, None, None, None, None, None,
]

_FILA_VALIDA_2 = [
    2, "2000000002", 987654321098765, date(2026, 7, 1), "PRORROGA",
    "15/01/2026", "15/02/2026", 32, 62, "G551 - G552",
    "1500000 - 1500000", "0 - 0", "17 - 15",
    date(2025, 6, 1), date(2026, 5, 1), 1200000,
    date(2026, 7, 5), "Aseguradora Test", "VIGENTE",
    None, None, None, None, None, None, None, None,
]

_FILA_MALFORMADA_FECHA_FINAL = [
    3, "3000000003", 555555555555555, date(2026, 8, 1), "INICIAL",
    "1/08/2026", "no-es-una-fecha", 20, 20, "J209",
    900000, 900000, 20,
    date(2025, 3, 1), date(2026, 2, 1), 400000,
    date(2026, 8, 5), "Aseguradora Test", "VIGENTE",
    None, None, None, None, None, None, None, None,
]

# Fila que parsea perfecto a nivel de campo (ningun parser de excel_common
# acota longitud de `radicado`) pero revienta a nivel de BD: la columna
# `IncapacidadPrevisional.radicado` es `String(20)` y este radicado trae 25
# digitos. Escrito como str python (no int/float) para que `parse_radicado`
# lo devuelva tal cual, sin normalizar longitud -- ver Finding 1, fix round 1.
_FILA_RADICADO_DEMASIADO_LARGO = [
    4, "4000000004", "9" * 25, date(2026, 9, 1), "INICIAL",
    "1/09/2026", "10/09/2026", 10, 10, "R101",
    900000, 900000, 10,
    date(2025, 4, 1), date(2026, 3, 1), 300000,
    date(2026, 9, 5), "Aseguradora Test", "VIGENTE",
    None, None, None, None, None, None, None, None,
]

# Par de filas para el encadenamiento de prorrogas (Finding 2, fix round 1):
# misma identificacion, la segunda (PRORROGA) empieza justo despues de que
# termina la primera (INICIAL).
_FILA_PRORROGA_INICIAL = [
    5, "5000000005", 111111111111111, date(2026, 1, 1), "INICIAL",
    "1/01/2026", "15/01/2026", 15, 15, "M545",
    1000000, 1000000, 15,
    date(2025, 1, 1), date(2025, 12, 1), 250000,
    date(2026, 1, 5), "Aseguradora Test", "VIGENTE",
    None, None, None, None, None, None, None, None,
]
_FILA_PRORROGA_SIGUIENTE = [
    6, "5000000005", 222222222222222, date(2026, 1, 16), "PRORROGA",
    "16/01/2026", "31/01/2026", 16, 31, "M545",
    1000000, 1000000, 16,
    date(2025, 1, 1), date(2025, 12, 1), 270000,
    date(2026, 1, 20), "Aseguradora Test", "VIGENTE",
    None, None, None, None, None, None, None, None,
]

# Par de filas repetidas (Finding 2, fix round 1): misma identificacion Y
# misma fecha_inicial -- una repetida genuina segun la regla AD de Task 1.4.
_FILA_REPETIDA_A = [
    7, "6000000006", 333333333333333, date(2026, 3, 1), "INICIAL",
    "1/03/2026", "10/03/2026", 10, 10, "M545",
    1000000, 1000000, 10,
    date(2025, 1, 1), date(2025, 12, 1), 200000,
    date(2026, 3, 5), "Aseguradora Test", "VIGENTE",
    None, None, None, None, None, None, None, None,
]
_FILA_REPETIDA_B = [
    8, "6000000006", 444444444444444, date(2026, 3, 1), "INICIAL",
    "1/03/2026", "10/03/2026", 10, 10, "M545",
    1000000, 1000000, 10,
    date(2025, 1, 1), date(2025, 12, 1), 200000,
    date(2026, 3, 5), "Aseguradora Test", "VIGENTE",
    None, None, None, None, None, None, None, None,
]


def _construir_workbook(filas: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "FORMALIZACION"
    for col_idx, valor in enumerate(_ENCABEZADO_FORMALIZACION, start=1):
        ws.cell(row=1, column=col_idx, value=valor)
    for row_idx, fila in enumerate(filas, start=2):
        for col_idx, valor in enumerate(fila, start=1):
            ws.cell(row=row_idx, column=col_idx, value=valor)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest_asyncio.fixture
async def smlmv_2026(db_session):
    """Siembra el SMLMV 2026 -- necesario porque los tests unit usan create_all, no migraciones."""
    smlmv = SmlmvParametros(ano=2026, valor=Decimal("1750905.00"), vigente_desde=date(2026, 1, 1))
    db_session.add(smlmv)
    await db_session.commit()
    await db_session.refresh(smlmv)
    return smlmv


# ---------------------------------------------------------------------------
# Lote con filas válidas
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_cargar_lote_con_filas_validas_persiste_incapacidades_y_periodos(
    db_session, test_usuario, smlmv_2026
):
    file_bytes = _construir_workbook([_FILA_VALIDA_1, _FILA_VALIDA_2])

    lote = await lote_previsional_service.cargar_lote(
        db_session,
        file_bytes=file_bytes,
        password=None,
        nombre_archivo="RADICADOS_TEST.xlsx",
        usuario_id=test_usuario.id,
    )

    assert lote.id is not None
    assert lote.total_filas == 2
    assert lote.total_incapacidades == 2
    # El libro original queda archivado en storage (Fase 6 lo necesita).
    assert lote.metadata_ is not None
    assert lote.metadata_.get("archivo_original_path")

    incapacidades = (
        (
            await db_session.execute(
                select(IncapacidadPrevisional).where(IncapacidadPrevisional.lote_id == lote.id)
            )
        )
        .scalars()
        .all()
    )
    assert len(incapacidades) == 2

    por_identificacion = {inc.identificacion: inc for inc in incapacidades}
    inc1 = por_identificacion["1000000001"]
    inc2 = por_identificacion["2000000002"]

    assert inc1.errores_carga is None
    assert inc1.cie10 == "M545"
    assert inc1.tipo_ingreso == "INICIAL"
    assert inc1.fecha_inicial == date(2026, 6, 1)
    assert inc1.fecha_final == date(2026, 6, 30)
    assert inc1.valor_auditado is not None
    assert inc1.radicado_normalizado == "0123456789012345"  # padding 15->16 dígitos

    assert inc2.errores_carga is None
    assert inc2.tipo_ingreso == "PRORROGA"

    periodos_inc1 = (
        (
            await db_session.execute(
                select(PeriodoPrevisional).where(PeriodoPrevisional.incapacidad_id == inc1.id)
            )
        )
        .scalars()
        .all()
    )
    assert len(periodos_inc1) == 1  # un solo mes completo

    periodos_inc2 = (
        (
            await db_session.execute(
                select(PeriodoPrevisional)
                .where(PeriodoPrevisional.incapacidad_id == inc2.id)
                .order_by(PeriodoPrevisional.orden)
            )
        )
        .scalars()
        .all()
    )
    assert len(periodos_inc2) == 2  # cruza frontera de mes (15/01 -> 15/02)
    assert periodos_inc2[0].dias == 17  # 15-31 enero
    assert periodos_inc2[1].dias == 15  # 1-15 febrero


# ---------------------------------------------------------------------------
# Lote con una fila malformada -- no aborta el lote completo
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_cargar_lote_fila_malformada_no_aborta_el_lote(db_session, test_usuario, smlmv_2026):
    file_bytes = _construir_workbook(
        [_FILA_VALIDA_1, _FILA_MALFORMADA_FECHA_FINAL, _FILA_VALIDA_2]
    )

    lote = await lote_previsional_service.cargar_lote(
        db_session,
        file_bytes=file_bytes,
        password=None,
        nombre_archivo="RADICADOS_CON_ERROR.xlsx",
        usuario_id=test_usuario.id,
    )

    assert lote.total_filas == 3
    assert lote.total_incapacidades == 2  # solo las 2 filas sin error cuentan como exitosas

    incapacidades = (
        (
            await db_session.execute(
                select(IncapacidadPrevisional).where(IncapacidadPrevisional.lote_id == lote.id)
            )
        )
        .scalars()
        .all()
    )
    assert len(incapacidades) == 3  # las 3 filas se persisten, incluida la mala

    por_identificacion = {inc.identificacion: inc for inc in incapacidades}
    inc_mala = por_identificacion["3000000003"]
    assert inc_mala.errores_carga is not None
    assert "fecha_final" in inc_mala.errores_carga

    periodos_inc_mala = (
        (
            await db_session.execute(
                select(PeriodoPrevisional).where(PeriodoPrevisional.incapacidad_id == inc_mala.id)
            )
        )
        .scalars()
        .all()
    )
    assert periodos_inc_mala == []  # sin segmentos: la fecha_final nunca se pudo parsear

    # Las otras 2 siguen bien, sin contaminarse por la fila mala.
    inc1 = por_identificacion["1000000001"]
    inc2 = por_identificacion["2000000002"]
    assert inc1.errores_carga is None
    assert inc2.errores_carga is None


# ---------------------------------------------------------------------------
# Rechazo de contraseña incorrecta a nivel de servicio (no solo del reader)
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_cargar_lote_password_incorrecta_rechaza_antes_de_escribir_en_bd(
    db_session, test_usuario
):
    file_bytes = FIXTURE_PATH.read_bytes()

    with pytest.raises(BadRequestException) as exc_info:
        await lote_previsional_service.cargar_lote(
            db_session,
            file_bytes=file_bytes,
            password="password-equivocada",
            nombre_archivo="afp_formalizacion_encrypted.xlsx",
            usuario_id=test_usuario.id,
        )

    assert "contraseña incorrecta" in str(exc_info.value)

    # Nada debe haber quedado persistido: el fallo es estructural, previo a
    # cualquier escritura en BD.
    lotes = (await db_session.execute(select(IncapacidadPrevisional))).scalars().all()
    assert lotes == []


# ---------------------------------------------------------------------------
# Encabezados críticos que no coinciden -- falla estructural, no aborta a mitad
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_cargar_lote_encabezado_critico_incorrecto_rechaza_sin_persistir(
    db_session, test_usuario
):
    from app.models.previsionales.lote_previsional import LotePrevisional

    encabezado_malo = list(_ENCABEZADO_FORMALIZACION)
    encabezado_malo[2] = "NUMERO RADICADO"  # columna C ("RADICADO") corrompida

    wb = Workbook()
    ws = wb.active
    ws.title = "FORMALIZACION"
    for col_idx, valor in enumerate(encabezado_malo, start=1):
        ws.cell(row=1, column=col_idx, value=valor)
    for col_idx, valor in enumerate(_FILA_VALIDA_1, start=1):
        ws.cell(row=2, column=col_idx, value=valor)
    buf = io.BytesIO()
    wb.save(buf)
    file_bytes = buf.getvalue()

    with pytest.raises(BadRequestException):
        await lote_previsional_service.cargar_lote(
            db_session,
            file_bytes=file_bytes,
            password=None,
            nombre_archivo="RADICADOS_ENCABEZADO_MALO.xlsx",
            usuario_id=test_usuario.id,
        )

    lotes = (await db_session.execute(select(LotePrevisional))).scalars().all()
    assert lotes == []


# ---------------------------------------------------------------------------
# Fix round 1 / Finding 1: un error de BD (no de parseo) en una fila no
# aborta el lote completo -- se persiste en modo degradado y las demas
# filas del lote siguen intactas.
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_cargar_lote_error_de_bd_en_una_fila_no_aborta_el_lote(
    db_session, test_usuario, smlmv_2026
):
    file_bytes = _construir_workbook(
        [_FILA_VALIDA_1, _FILA_RADICADO_DEMASIADO_LARGO, _FILA_VALIDA_2]
    )

    # (a) `cargar_lote` no debe propagar el `DataError` de BD.
    lote = await lote_previsional_service.cargar_lote(
        db_session,
        file_bytes=file_bytes,
        password=None,
        nombre_archivo="RADICADOS_RADICADO_LARGO.xlsx",
        usuario_id=test_usuario.id,
    )

    assert lote.total_filas == 3
    # La fila con error de BD no cuenta como "creada exitosamente".
    assert lote.total_incapacidades == 2

    incapacidades = (
        (
            await db_session.execute(
                select(IncapacidadPrevisional).where(IncapacidadPrevisional.lote_id == lote.id)
            )
        )
        .scalars()
        .all()
    )
    # (b) las 2 filas válidas SÍ se persisten y (c) la fila mala tambien
    # queda persistida (modo degradado) -- ninguna de las 3 desaparece.
    assert len(incapacidades) == 3

    por_identificacion = {inc.identificacion: inc for inc in incapacidades}

    inc_bd_error = por_identificacion["4000000004"]
    assert inc_bd_error.errores_carga is not None
    assert "persistencia" in inc_bd_error.errores_carga
    # El radicado de 25 digitos se truncó a la longitud máxima de columna
    # SOLO en este camino degradado -- nunca en el camino feliz.
    assert len(inc_bd_error.radicado) <= 20
    assert len(inc_bd_error.radicado_normalizado) <= 16

    # Los periodos se omiten deliberadamente en el reintento degradado.
    periodos_fila_mala = (
        (
            await db_session.execute(
                select(PeriodoPrevisional).where(
                    PeriodoPrevisional.incapacidad_id == inc_bd_error.id
                )
            )
        )
        .scalars()
        .all()
    )
    assert periodos_fila_mala == []

    # Las filas válidas no se contaminan por el fallo de BD de la otra fila.
    inc1 = por_identificacion["1000000001"]
    inc2 = por_identificacion["2000000002"]
    assert inc1.errores_carga is None
    assert inc2.errores_carga is None
    assert inc1.radicado == "123456789012345"

    periodos_inc1 = (
        (
            await db_session.execute(
                select(PeriodoPrevisional).where(PeriodoPrevisional.incapacidad_id == inc1.id)
            )
        )
        .scalars()
        .all()
    )
    assert len(periodos_inc1) == 1


# ---------------------------------------------------------------------------
# Fix round 1 / Finding 2: cableado de `agrupar_repetidas`/`encadenar_prorrogas`
# hacia `es_duplicado_interno`/`prorroga_de_id` -- sin ejercitar antes.
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_cargar_lote_encadena_prorroga_de_id_con_la_incapacidad_anterior(
    db_session, test_usuario, smlmv_2026
):
    file_bytes = _construir_workbook([_FILA_PRORROGA_INICIAL, _FILA_PRORROGA_SIGUIENTE])

    lote = await lote_previsional_service.cargar_lote(
        db_session,
        file_bytes=file_bytes,
        password=None,
        nombre_archivo="RADICADOS_PRORROGA.xlsx",
        usuario_id=test_usuario.id,
    )

    incapacidades = (
        (
            await db_session.execute(
                select(IncapacidadPrevisional).where(IncapacidadPrevisional.lote_id == lote.id)
            )
        )
        .scalars()
        .all()
    )
    assert len(incapacidades) == 2

    por_tipo = {inc.tipo_ingreso: inc for inc in incapacidades}
    inicial = por_tipo["INICIAL"]
    prorroga = por_tipo["PRORROGA"]

    assert inicial.prorroga_de_id is None
    assert prorroga.prorroga_de_id == inicial.id


@pytest.mark.asyncio
async def test_cargar_lote_marca_es_duplicado_interno_en_filas_repetidas(
    db_session, test_usuario, smlmv_2026
):
    file_bytes = _construir_workbook([_FILA_REPETIDA_A, _FILA_REPETIDA_B])

    lote = await lote_previsional_service.cargar_lote(
        db_session,
        file_bytes=file_bytes,
        password=None,
        nombre_archivo="RADICADOS_REPETIDAS.xlsx",
        usuario_id=test_usuario.id,
    )

    incapacidades = (
        (
            await db_session.execute(
                select(IncapacidadPrevisional).where(IncapacidadPrevisional.lote_id == lote.id)
            )
        )
        .scalars()
        .all()
    )
    assert len(incapacidades) == 2

    # Misma identificacion + misma fecha_inicial -> repetida genuina (regla
    # AD): AMBAS filas quedan marcadas, simetricamente, sin elegir una
    # "original" (ver comentario junto al cableado en lote_service.py).
    assert all(inc.es_duplicado_interno for inc in incapacidades)
    # `incapacidad_origen_id` deliberadamente no se pobla por esta carga
    # (no hay campo "esta es la copia de aquella" que el diseño llene aqui).
    assert all(inc.incapacidad_origen_id is None for inc in incapacidades)
