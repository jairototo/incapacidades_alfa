"""GET /empleados/plantilla: downloadable .xlsx with the right headers, no banking columns."""
import io

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook


@pytest.mark.asyncio
async def test_descargar_plantilla_empleados(client: AsyncClient, admin_token_headers):
    resp = await client.get("/api/v1/empleados/plantilla", headers=admin_token_headers)
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb["Empleados"]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]

    assert headers == [
        "numero_documento", "tipo_documento", "nombres", "apellidos", "email",
        "telefono", "fecha_nacimiento", "genero", "cargo", "area",
        "fecha_ingreso", "salario_base", "nit_empresa",
    ]
    for banking_field in ("cuenta_bancaria", "banco", "tipo_cuenta"):
        assert banking_field not in headers

    assert "Instrucciones" in wb.sheetnames


@pytest.mark.asyncio
async def test_descargar_plantilla_forbidden_for_liquidador(client: AsyncClient, db_session):
    from app.core.security import create_access_token, get_password_hash
    from app.models.usuario import Usuario
    from app.utils.enums import RolUsuario, EstadoUsuario

    usuario = Usuario(
        username="liquidador_plantilla",
        email="liquidador_plantilla@example.com",
        password_hash=get_password_hash("Test1234"),
        nombre_completo="Test LIQUIDADOR",
        rol=RolUsuario.LIQUIDADOR,
        estado=EstadoUsuario.ACTIVO,
    )
    db_session.add(usuario)
    await db_session.commit()
    await db_session.refresh(usuario)
    token = create_access_token(data={"sub": str(usuario.id), "token_version": usuario.token_version})

    resp = await client.get(
        "/api/v1/empleados/plantilla", headers={"Authorization": f"Bearer {token}"}
    )
    assert resp.status_code == 403
