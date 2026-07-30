"""
Tests para el adaptador de datos de referencia previsionales (Task 3.2,
`app/services/previsionales/referencia_adapter.py`).

Construye workbooks openpyxl sintéticos en memoria, con el layout de
encabezados REAL verificado esta sesión contra el workbook de producción
`RADICADOS_AUDITORIA_20260706.xlsx` (hojas SINIESTROS, SOLICITUDES,
"LISTADO ITE DIA") — no se usa el archivo real de 9 MB para los tests
principales (demasiado lento para tests unitarios). Un test opcional al
final sí usa el fixture real recortado como smoke test.

Cobertura:
- Import válido de un par de filas por hoja.
- Una fila malformada no aborta el resto de la hoja.
- Validación de encabezados rechaza una hoja con forma incorrecta.
- Deduplicación/idempotencia en re-importación (mismo archivo dos veces,
  y también contra filas ya existentes en BD antes de importar).
"""
import io
from datetime import date

import openpyxl
import pytest
from sqlalchemy import select

from app.core.exceptions import BadRequestException
from app.models.previsionales.ite_historico import IteHistorico
from app.models.previsionales.siniestro_previsional import SiniestroPrevisional
from app.models.previsionales.solicitud_previsional import SolicitudPrevisional
from app.services.previsionales.referencia_adapter import (
    ExcelReferenciaAdapter,
    excel_referencia_adapter,
)

# NOTA: no se declara `pytestmark = pytest.mark.asyncio` -- `pytest.ini`
# tiene `asyncio_mode = auto` (ver CLAUDE.md), que detecta automáticamente
# las funciones `async def` como tests asyncio sin necesidad del marcador
# explícito. Declararlo a nivel de módulo marcaría también al único test
# síncrono (`test_singleton_expuesto`), generando una advertencia de pytest.


# ---------------------------------------------------------------------------
# Helpers para construir workbooks sintéticos
# ---------------------------------------------------------------------------


def _wb_bytes(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _set_row(ws, row_num: int, values: dict[str, object]) -> None:
    for letter, value in values.items():
        ws[f"{letter}{row_num}"] = value


def _wb_siniestros(filas: list[dict[str, object]], header: dict[str, str] | None = None) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SINIESTROS"
    _set_row(
        ws,
        1,
        header
        or {
            "A": "Número Id",
            "B": "No. Siniestro",
            "E": "Fecha Siniestro",
            "G": "Estado siniestro",
            "S": "Origen siniestro",
            "T": "Fecha Aviso",
        },
    )
    for i, fila in enumerate(filas, start=2):
        _set_row(ws, i, fila)
    return _wb_bytes(wb)


def _wb_solicitudes(filas: list[dict[str, object]], header: dict[str, str] | None = None) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SOLICITUDES"
    _set_row(
        ws,
        1,
        header
        or {
            "B": "No. Solicitud Auditoria",
            "D": "Número Id",
            "N": "Dia 181",
            "R": "Observaciones",
            "AF": "FECHA_CRIE",
        },
    )
    for i, fila in enumerate(filas, start=2):
        _set_row(ws, i, fila)
    return _wb_bytes(wb)


def _wb_ite(filas: list[dict[str, object]], header: dict[str, str] | None = None) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "LISTADO ITE DIA"
    _set_row(
        ws,
        1,
        header
        or {
            "F": "Número identificacion",
            "I": "Fecha inicial",
            "J": "Fecha final",
        },
    )
    for i, fila in enumerate(filas, start=2):
        _set_row(ws, i, fila)
    return _wb_bytes(wb)


# ---------------------------------------------------------------------------
# SINIESTROS
# ---------------------------------------------------------------------------


async def test_importar_siniestros_filas_validas(db_session):
    file_bytes = _wb_siniestros(
        [
            {
                "A": "1000000001",
                "B": "SIN-001",
                "E": date(2026, 1, 1),
                "G": "AVISO",
                "S": "COMUN",
                "T": date(2026, 1, 2),
            },
            {
                "A": "1000000002",
                "B": "SIN-002",
                "E": date(2026, 2, 1),
                "G": "ABIERTO",
                "S": "LABORAL",
                "T": date(2026, 2, 3),
            },
        ]
    )

    adapter = ExcelReferenciaAdapter()
    count = await adapter.importar_siniestros(db_session, file_bytes)

    assert count == 2

    result = await db_session.execute(select(SiniestroPrevisional))
    rows = {r.numero_siniestro: r for r in result.scalars().all()}
    assert set(rows) == {"SIN-001", "SIN-002"}
    assert rows["SIN-001"].identificacion == "1000000001"
    assert rows["SIN-001"].estado == "AVISO"
    assert rows["SIN-001"].origen == "COMUN"
    assert rows["SIN-001"].fecha_siniestro == date(2026, 1, 1)
    assert rows["SIN-001"].fecha_aviso == date(2026, 1, 2)


async def test_importar_siniestros_fila_malformada_no_aborta_las_demas(db_session):
    file_bytes = _wb_siniestros(
        [
            {
                "A": "1000000001",
                "B": "SIN-001",
                "E": "fecha-invalida-no-parseable",
                "G": "AVISO",
                "S": "COMUN",
                "T": date(2026, 1, 2),
            },
            {
                "A": "1000000002",
                "B": "SIN-002",
                "E": date(2026, 2, 1),
                "G": "ABIERTO",
                "S": "LABORAL",
                "T": date(2026, 2, 3),
            },
        ]
    )

    adapter = ExcelReferenciaAdapter()
    count = await adapter.importar_siniestros(db_session, file_bytes)

    # La fila 1 (fecha inválida) se omite; la fila 2 se importa igual.
    assert count == 1
    result = await db_session.execute(select(SiniestroPrevisional))
    rows = result.scalars().all()
    assert len(rows) == 1
    assert rows[0].numero_siniestro == "SIN-002"


async def test_importar_siniestros_fila_sin_identificacion_se_omite(db_session):
    file_bytes = _wb_siniestros(
        [
            {"A": None, "B": "SIN-001", "E": date(2026, 1, 1), "G": "AVISO", "S": "COMUN", "T": date(2026, 1, 2)},
            {"A": "1000000002", "B": "SIN-002", "E": date(2026, 2, 1), "G": "ABIERTO", "S": "LABORAL", "T": date(2026, 2, 3)},
        ]
    )

    adapter = ExcelReferenciaAdapter()
    count = await adapter.importar_siniestros(db_session, file_bytes)

    assert count == 1


async def test_importar_siniestros_header_invalido_rechaza(db_session):
    file_bytes = _wb_siniestros(
        [{"A": "1000000001", "B": "SIN-001", "E": date(2026, 1, 1), "G": "AVISO", "S": "COMUN", "T": date(2026, 1, 2)}],
        header={
            "A": "Columna Incorrecta",
            "B": "No. Siniestro",
            "E": "Fecha Siniestro",
            "G": "Estado siniestro",
            "S": "Origen siniestro",
            "T": "Fecha Aviso",
        },
    )

    adapter = ExcelReferenciaAdapter()
    with pytest.raises(BadRequestException):
        await adapter.importar_siniestros(db_session, file_bytes)


async def test_importar_siniestros_reimportar_mismo_archivo_es_idempotente(db_session):
    file_bytes = _wb_siniestros(
        [{"A": "1000000001", "B": "SIN-001", "E": date(2026, 1, 1), "G": "AVISO", "S": "COMUN", "T": date(2026, 1, 2)}]
    )

    adapter = ExcelReferenciaAdapter()
    primera = await adapter.importar_siniestros(db_session, file_bytes)
    segunda = await adapter.importar_siniestros(db_session, file_bytes)

    assert primera == 1
    assert segunda == 0  # misma llave (identificacion, numero_siniestro) -- no duplica

    result = await db_session.execute(select(SiniestroPrevisional))
    assert len(result.scalars().all()) == 1


# ---------------------------------------------------------------------------
# SOLICITUDES
# ---------------------------------------------------------------------------


async def test_importar_solicitudes_filas_validas(db_session):
    file_bytes = _wb_solicitudes(
        [
            {
                "B": "RAD-001",
                "D": "2000000001",
                "N": date(2026, 3, 1),
                "R": "observación de prueba",
                "AF": date(2026, 3, 15),
            },
            {
                "B": "RAD-002",
                "D": "2000000002",
                "N": date(2026, 4, 1),
                "R": None,
                "AF": None,
            },
        ]
    )

    adapter = ExcelReferenciaAdapter()
    count = await adapter.importar_solicitudes(db_session, file_bytes)

    assert count == 2
    result = await db_session.execute(select(SolicitudPrevisional))
    rows = {r.radicado: r for r in result.scalars().all()}
    assert set(rows) == {"RAD-001", "RAD-002"}
    assert rows["RAD-001"].identificacion == "2000000001"
    assert rows["RAD-001"].dia_181 == date(2026, 3, 1)
    assert rows["RAD-001"].fecha_crie == date(2026, 3, 15)
    assert rows["RAD-001"].observacion == "observación de prueba"
    # tipo_ingreso/fecha_inicial/fecha_final: GAP documentado, sin mapeo de origen.
    assert rows["RAD-001"].tipo_ingreso is None
    assert rows["RAD-001"].fecha_inicial is None


async def test_importar_solicitudes_fila_malformada_no_aborta_las_demas(db_session):
    file_bytes = _wb_solicitudes(
        [
            {"B": "RAD-001", "D": "2000000001", "N": "no-es-fecha", "R": None, "AF": None},
            {"B": "RAD-002", "D": "2000000002", "N": date(2026, 4, 1), "R": None, "AF": None},
        ]
    )

    adapter = ExcelReferenciaAdapter()
    count = await adapter.importar_solicitudes(db_session, file_bytes)

    assert count == 1
    result = await db_session.execute(select(SolicitudPrevisional))
    rows = result.scalars().all()
    assert len(rows) == 1
    assert rows[0].radicado == "RAD-002"


async def test_importar_solicitudes_header_invalido_rechaza(db_session):
    file_bytes = _wb_solicitudes(
        [{"B": "RAD-001", "D": "2000000001", "N": date(2026, 3, 1), "R": None, "AF": None}],
        header={
            "B": "No. Solicitud Auditoria",
            "D": "Columna Incorrecta",
            "N": "Dia 181",
            "R": "Observaciones",
            "AF": "FECHA_CRIE",
        },
    )

    adapter = ExcelReferenciaAdapter()
    with pytest.raises(BadRequestException):
        await adapter.importar_solicitudes(db_session, file_bytes)


async def test_importar_solicitudes_sin_radicado_siempre_se_inserta(db_session):
    """Sin `radicado` no hay llave natural para deduplicar -- cada import inserta la fila."""
    file_bytes = _wb_solicitudes(
        [{"B": None, "D": "2000000001", "N": date(2026, 3, 1), "R": None, "AF": None}]
    )

    adapter = ExcelReferenciaAdapter()
    primera = await adapter.importar_solicitudes(db_session, file_bytes)
    segunda = await adapter.importar_solicitudes(db_session, file_bytes)

    assert primera == 1
    assert segunda == 1  # no deduplicable sin radicado -- documentado en el módulo

    result = await db_session.execute(select(SolicitudPrevisional))
    assert len(result.scalars().all()) == 2


# ---------------------------------------------------------------------------
# LISTADO ITE DIA
# ---------------------------------------------------------------------------


async def test_importar_ite_historico_filas_validas(db_session):
    file_bytes = _wb_ite(
        [
            {"F": "3000000001", "I": date(2026, 5, 1), "J": date(2026, 5, 10)},
            {"F": "3000000002", "I": date(2026, 6, 1), "J": None},
        ]
    )

    adapter = ExcelReferenciaAdapter()
    count = await adapter.importar_ite_historico(db_session, file_bytes)

    assert count == 2
    result = await db_session.execute(select(IteHistorico))
    rows = {r.identificacion: r for r in result.scalars().all()}
    assert rows["3000000001"].fecha_inicial == date(2026, 5, 1)
    assert rows["3000000001"].fecha_final == date(2026, 5, 10)
    assert rows["3000000002"].fecha_final is None
    # valor_pagado/fecha_pago: GAP documentado, sin mapeo de origen inequívoco.
    assert rows["3000000001"].valor_pagado is None
    assert rows["3000000001"].fecha_pago is None


async def test_importar_ite_historico_fila_malformada_no_aborta_las_demas(db_session):
    file_bytes = _wb_ite(
        [
            {"F": "3000000001", "I": "fecha-invalida", "J": None},
            {"F": "3000000002", "I": date(2026, 6, 1), "J": None},
        ]
    )

    adapter = ExcelReferenciaAdapter()
    count = await adapter.importar_ite_historico(db_session, file_bytes)

    assert count == 1
    result = await db_session.execute(select(IteHistorico))
    rows = result.scalars().all()
    assert len(rows) == 1
    assert rows[0].identificacion == "3000000002"


async def test_importar_ite_historico_fila_sin_fecha_inicial_se_omite(db_session):
    file_bytes = _wb_ite(
        [
            {"F": "3000000001", "I": None, "J": None},
            {"F": "3000000002", "I": date(2026, 6, 1), "J": None},
        ]
    )

    adapter = ExcelReferenciaAdapter()
    count = await adapter.importar_ite_historico(db_session, file_bytes)

    assert count == 1


async def test_importar_ite_historico_header_invalido_rechaza(db_session):
    file_bytes = _wb_ite(
        [{"F": "3000000001", "I": date(2026, 5, 1), "J": None}],
        header={"F": "Columna Incorrecta", "I": "Fecha inicial", "J": "Fecha final"},
    )

    adapter = ExcelReferenciaAdapter()
    with pytest.raises(BadRequestException):
        await adapter.importar_ite_historico(db_session, file_bytes)


async def test_importar_ite_historico_omite_clave_ya_existente_en_bd(db_session):
    """
    La regla AT (doble pago) depende de que `ite_historico` refleje pagos ya
    hechos -- si ya existe `(identificacion, fecha_inicial)` en BD (p.ej.
    sembrado por una importación anterior o por otra fuente), la
    reimportación NO debe duplicar esa fila.
    """
    existente = IteHistorico(identificacion="3000000001", fecha_inicial=date(2026, 5, 1))
    db_session.add(existente)
    await db_session.commit()

    file_bytes = _wb_ite(
        [
            {"F": "3000000001", "I": date(2026, 5, 1), "J": date(2026, 5, 10)},
            {"F": "3000000002", "I": date(2026, 6, 1), "J": None},
        ]
    )

    adapter = ExcelReferenciaAdapter()
    count = await adapter.importar_ite_historico(db_session, file_bytes)

    assert count == 1  # solo la fila 3000000002 es nueva
    result = await db_session.execute(select(IteHistorico))
    rows = result.scalars().all()
    assert len(rows) == 2  # la sembrada + la nueva, sin duplicar la existente


async def test_importar_ite_historico_reimportar_mismo_archivo_es_idempotente(db_session):
    file_bytes = _wb_ite([{"F": "3000000001", "I": date(2026, 5, 1), "J": date(2026, 5, 10)}])

    adapter = ExcelReferenciaAdapter()
    primera = await adapter.importar_ite_historico(db_session, file_bytes)
    segunda = await adapter.importar_ite_historico(db_session, file_bytes)

    assert primera == 1
    assert segunda == 0

    result = await db_session.execute(select(IteHistorico))
    assert len(result.scalars().all()) == 1


# ---------------------------------------------------------------------------
# Errores estructurales / hoja faltante
# ---------------------------------------------------------------------------


async def test_importar_siniestros_archivo_no_es_excel_valido(db_session):
    adapter = ExcelReferenciaAdapter()
    with pytest.raises(BadRequestException):
        await adapter.importar_siniestros(db_session, b"esto no es un xlsx")


async def test_importar_siniestros_hoja_faltante(db_session):
    wb = openpyxl.Workbook()
    wb.active.title = "OTRA_HOJA"
    file_bytes = _wb_bytes(wb)

    adapter = ExcelReferenciaAdapter()
    with pytest.raises(BadRequestException):
        await adapter.importar_siniestros(db_session, file_bytes)


def test_singleton_expuesto():
    """El módulo expone `excel_referencia_adapter` como singleton, mismo
    patrón que el resto de servicios previsionales."""
    assert isinstance(excel_referencia_adapter, ExcelReferenciaAdapter)


# ---------------------------------------------------------------------------
# Smoke test OPCIONAL contra el workbook real recortado (Task 3.1 fixture)
#
# NOTA: esto es más un smoke test de integración que un test unitario --
# usa el archivo real `Copy_RADICADOS_AUDITORIA_20260706.xlsx` (recortado,
# ~300KB, confirmado NO cifrado) en vez de un workbook sintético, para
# probar el adaptador contra datos con forma de producción genuina. No es
# necesario para la cobertura principal (ya cubierta arriba con workbooks
# sintéticos) -- se incluye como validación adicional de que el mapeo de
# columnas verificado esta sesión sigue siendo válido contra el archivo real.
# ---------------------------------------------------------------------------

from pathlib import Path  # noqa: E402

_FIXTURE_REAL = Path(__file__).parent.parent / "fixtures" / "Copy_RADICADOS_AUDITORIA_20260706.xlsx"


@pytest.mark.skipif(not _FIXTURE_REAL.exists(), reason="fixture real no disponible")
async def test_smoke_importar_desde_workbook_real_recortado(db_session):
    file_bytes = _FIXTURE_REAL.read_bytes()

    adapter = ExcelReferenciaAdapter()
    n_siniestros = await adapter.importar_siniestros(db_session, file_bytes)
    n_solicitudes = await adapter.importar_solicitudes(db_session, file_bytes)
    n_ite = await adapter.importar_ite_historico(db_session, file_bytes)

    # El archivo recortado trae filas reales en las 3 hojas (ver reporte de
    # Task 3.2) -- solo se afirma que hay AL MENOS una fila importada por
    # hoja, sin acoplarse a un conteo exacto que pueda cambiar si el
    # fixture se regenera.
    assert n_siniestros > 0
    assert n_solicitudes > 0
    assert n_ite > 0
