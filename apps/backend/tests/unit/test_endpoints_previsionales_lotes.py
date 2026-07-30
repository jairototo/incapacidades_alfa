"""
Tests para los endpoints REST del módulo Previsionales (Task 4.1):
lotes.py, referencia.py, parametros.py.

Usa `client` (httpx AsyncClient contra la app real, `get_db` sobreescrito
con `db_session` -- ver tests/conftest.py) y libros sintéticos en memoria
con la forma exacta de la hoja FORMALIZACION (columnas A-AA), tal como
`test_lote_service.py` (Task 3.1) ya verificó contra el workbook real de
producción.

Nota sobre el fixture cifrado (tests/fixtures/afp_formalizacion_encrypted.xlsx,
password "test1234", Task 0.3): sus encabezados
(TIPO_DOC/NUMERO_DOC/NOMBRE_AFILIADO/MONTO) son un proof-of-concept
genérico de que msoffcrypto puede desencriptar dentro de este repo -- NO
coinciden con los encabezados críticos reales de FORMALIZACION
(# IDENTIFICACION, RADICADO, etc.), así que un "happy path" completo con
contraseña CORRECTA fallaría igual en `validar_encabezados` sin que eso
pruebe nada sobre el endpoint. Por eso, igual que `test_lote_service.py`,
este archivo usa ese fixture solo para el caso "contraseña incorrecta"
(falla antes de llegar a validar encabezados) y usa un workbook sintético
sin cifrar para el happy path -- que sí ejercita la ruta de negocio real
end-to-end (parseo -> segmentación -> liquidación -> persistencia).
"""
from __future__ import annotations

import io
from datetime import date
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest
import pytest_asyncio
from openpyxl import Workbook

from app.core.security import create_access_token, get_password_hash
from app.models.previsionales.smlmv_parametros import SmlmvParametros
from app.models.usuario import Usuario
from app.utils.enums import EstadoUsuario, RolUsuario

FIXTURE_PATH = Path(__file__).parent.parent / "fixtures" / "afp_formalizacion_encrypted.xlsx"
FIXTURE_PASSWORD = "test1234"

_XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_ENCABEZADO_FORMALIZACION = [
    "Junio",
    "# IDENTIFICACION",
    "RADICADO",
    "FECHA RADICACION AFP",
    "TIPO INGRESO",
    "FECHA_INICIAL",
    "FECHA_FINAL",
    "No. DIAS",
    "DIAS ACUMULADOS",
    "DIAGNOSTICOS",
    "IBC",
    "SALARIO",
    "DIAS COTIZADOS",
    "FECHA DIA 181",
    "FECHA DIA 360",
    "VALOR",
    "FECHA RADICACION ASEGURADORA",
    "ASEGURADORA",
    "ESTADO AFILIADO",
    "DIA 181 ALFA",
    "FECHA CRIE ALFA",
    "OBSERVACIÓN ALFA",
    "OBSERVACIÓN PORVENIR",
    "VALOR PAGADO",
    "FECHA DE PAGO ASEGURADORA",
    "INSTANCIA JUDICIAL (cuando aplique)",
    "OBSERVACION JURIDICA (cuando aplique)",
]

_FILA_VALIDA_1 = [
    1, "1000000001", 123456789012345, date(2026, 6, 1), "INICIAL",
    "1/06/2026", "30/06/2026", 30, 30, "M545",
    1000000, 1000000, 30,
    date(2025, 1, 1), date(2025, 12, 1), 500000,
    date(2026, 6, 5), "Aseguradora Test", "VIGENTE",
    None, None, None, None, None, None, None, None,
]

_FILA_MALFORMADA = [
    3, "3000000003", 555555555555555, date(2026, 8, 1), "INICIAL",
    "1/08/2026", "no-es-una-fecha", 20, 20, "J209",
    900000, 900000, 20,
    date(2025, 3, 1), date(2026, 2, 1), 400000,
    date(2026, 8, 5), "Aseguradora Test", "VIGENTE",
    None, None, None, None, None, None, None, None,
]


def _construir_workbook_formalizacion(filas: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "FORMALIZACION"
    for col_idx, valor in enumerate(_ENCABEZADO_FORMALIZACION, start=1):
        ws.cell(row=1, column=col_idx, value=valor)
    for row_idx, fila in enumerate(filas, start=2):
        for col_idx, valor in enumerate(fila, start=1):
            ws.cell(row=row_idx, column=col_idx, value=valor)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _construir_workbook_siniestros(filas: list[list]) -> bytes:
    """Solo las columnas que `referencia_adapter._ENCABEZADOS_SINIESTROS`
    valida (A, B, E, G, S, T) importan; el resto puede quedar vacío."""
    encabezado = [""] * 20
    encabezado[0] = "Número Id"
    encabezado[1] = "No. Siniestro"
    encabezado[4] = "Fecha Siniestro"
    encabezado[6] = "Estado siniestro"
    encabezado[18] = "Origen siniestro"
    encabezado[19] = "Fecha Aviso"

    wb = Workbook()
    ws = wb.active
    ws.title = "SINIESTROS"
    for col_idx, valor in enumerate(encabezado, start=1):
        ws.cell(row=1, column=col_idx, value=valor)
    for row_idx, fila in enumerate(filas, start=2):
        for col_idx, valor in enumerate(fila, start=1):
            ws.cell(row=row_idx, column=col_idx, value=valor)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------


@pytest_asyncio.fixture
async def smlmv_2026(db_session):
    """Siembra el SMLMV 2026 -- necesario porque los tests unit usan
    create_all (no migraciones), la tabla smlmv_parametros arranca vacía y
    `cargar_lote` levanta ValueError si no hay SMLMV configurado para el
    año de un segmento."""
    smlmv = SmlmvParametros(ano=2026, valor=Decimal("1750905.00"), vigente_desde=date(2026, 1, 1))
    db_session.add(smlmv)
    await db_session.commit()
    await db_session.refresh(smlmv)
    return smlmv


async def _crear_usuario_con_rol(db_session, rol: RolUsuario, username: str) -> Usuario:
    usuario = Usuario(
        username=username,
        email=f"{username}@example.com",
        password_hash=get_password_hash("Test1234!"),
        nombre_completo=f"Usuario {rol.value}",
        rol=rol,
        estado=EstadoUsuario.ACTIVO,
    )
    db_session.add(usuario)
    await db_session.commit()
    await db_session.refresh(usuario)
    return usuario


def _headers_para(usuario: Usuario) -> dict:
    token = create_access_token(
        data={"sub": str(usuario.id), "token_version": usuario.token_version}
    )
    return {"Authorization": f"Bearer {token}"}


@pytest_asyncio.fixture
async def auditor_previsionales_headers(db_session) -> dict:
    usuario = await _crear_usuario_con_rol(
        db_session, RolUsuario.AUDITOR_PREVISIONALES, "auditor_previsionales_ep"
    )
    return _headers_para(usuario)


@pytest_asyncio.fixture
async def readonly_headers(db_session) -> dict:
    usuario = await _crear_usuario_con_rol(db_session, RolUsuario.READONLY, "readonly_previsionales_ep")
    return _headers_para(usuario)


# ---------------------------------------------------------------------------
# POST /previsionales/lotes
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_post_lotes_happy_path_crea_lote_y_persiste_incapacidades(
    client, auditor_previsionales_headers, smlmv_2026
):
    file_bytes = _construir_workbook_formalizacion([_FILA_VALIDA_1])

    resp = await client.post(
        "/api/v1/previsionales/lotes",
        headers=auditor_previsionales_headers,
        files={"file": ("RADICADOS_TEST.xlsx", file_bytes, _XLSX_CONTENT_TYPE)},
        data={"nombre_archivo": "RADICADOS_TEST.xlsx"},
    )

    assert resp.status_code == 201, resp.text
    data = resp.json()
    assert data["nombre_archivo"] == "RADICADOS_TEST.xlsx"
    assert data["total_filas"] == 1
    assert data["total_incapacidades"] == 1
    assert data["id"] is not None


@pytest.mark.asyncio
async def test_post_lotes_sin_nombre_archivo_usa_el_nombre_del_upload(
    client, auditor_previsionales_headers, smlmv_2026
):
    file_bytes = _construir_workbook_formalizacion([_FILA_VALIDA_1])

    resp = await client.post(
        "/api/v1/previsionales/lotes",
        headers=auditor_previsionales_headers,
        files={"file": ("desde_el_upload.xlsx", file_bytes, _XLSX_CONTENT_TYPE)},
    )

    assert resp.status_code == 201, resp.text
    assert resp.json()["nombre_archivo"] == "desde_el_upload.xlsx"


@pytest.mark.asyncio
async def test_post_lotes_password_incorrecta_en_archivo_cifrado_devuelve_400(
    client, auditor_previsionales_headers
):
    """Usa el fixture cifrado real de Task 0.3 (afp_formalizacion_encrypted.xlsx)
    con una contraseña incorrecta -- debe rechazarse ANTES de llegar a
    validar encabezados de negocio (ver nota del módulo)."""
    file_bytes = FIXTURE_PATH.read_bytes()

    resp = await client.post(
        "/api/v1/previsionales/lotes",
        headers=auditor_previsionales_headers,
        files={"file": ("afp_formalizacion_encrypted.xlsx", file_bytes, _XLSX_CONTENT_TYPE)},
        data={"password": "password-equivocada"},
    )

    assert resp.status_code == 400, resp.text
    assert "contraseña incorrecta" in resp.json()["detail"]


@pytest.mark.asyncio
async def test_post_lotes_archivo_cifrado_sin_password_devuelve_400(
    client, auditor_previsionales_headers
):
    """`password` es opcional en el request, pero si el archivo SÍ está
    cifrado y no se manda, debe rechazarse con 400 (no 500)."""
    file_bytes = FIXTURE_PATH.read_bytes()

    resp = await client.post(
        "/api/v1/previsionales/lotes",
        headers=auditor_previsionales_headers,
        files={"file": ("afp_formalizacion_encrypted.xlsx", file_bytes, _XLSX_CONTENT_TYPE)},
    )

    assert resp.status_code == 400, resp.text
    assert "requiere contraseña" in resp.json()["detail"]


@pytest.mark.asyncio
async def test_post_lotes_sin_permiso_devuelve_403(client, readonly_headers):
    file_bytes = _construir_workbook_formalizacion([_FILA_VALIDA_1])

    resp = await client.post(
        "/api/v1/previsionales/lotes",
        headers=readonly_headers,
        files={"file": ("RADICADOS_TEST.xlsx", file_bytes, _XLSX_CONTENT_TYPE)},
    )

    assert resp.status_code == 403


@pytest.mark.asyncio
async def test_post_lotes_sin_token_devuelve_401(client):
    file_bytes = _construir_workbook_formalizacion([_FILA_VALIDA_1])

    resp = await client.post(
        "/api/v1/previsionales/lotes",
        files={"file": ("RADICADOS_TEST.xlsx", file_bytes, _XLSX_CONTENT_TYPE)},
    )

    assert resp.status_code == 401


# ---------------------------------------------------------------------------
# GET /previsionales/lotes  (lista + contadores)
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_get_lotes_lista_incluye_el_lote_recien_creado(
    client, auditor_previsionales_headers, smlmv_2026
):
    file_bytes = _construir_workbook_formalizacion([_FILA_VALIDA_1])
    post_resp = await client.post(
        "/api/v1/previsionales/lotes",
        headers=auditor_previsionales_headers,
        files={"file": ("LOTE_LISTA.xlsx", file_bytes, _XLSX_CONTENT_TYPE)},
    )
    lote_id = post_resp.json()["id"]

    resp = await client.get("/api/v1/previsionales/lotes", headers=auditor_previsionales_headers)

    assert resp.status_code == 200, resp.text
    ids = [item["id"] for item in resp.json()]
    assert lote_id in ids
    encontrado = next(item for item in resp.json() if item["id"] == lote_id)
    assert encontrado["total_incapacidades"] == 1
    assert encontrado["total_filas"] == 1


@pytest.mark.asyncio
async def test_get_lotes_respeta_paginacion_skip_limit(
    client, auditor_previsionales_headers, smlmv_2026
):
    for i in range(3):
        file_bytes = _construir_workbook_formalizacion([_FILA_VALIDA_1])
        await client.post(
            "/api/v1/previsionales/lotes",
            headers=auditor_previsionales_headers,
            files={"file": (f"LOTE_PAG_{i}.xlsx", file_bytes, _XLSX_CONTENT_TYPE)},
        )

    resp = await client.get(
        "/api/v1/previsionales/lotes",
        headers=auditor_previsionales_headers,
        params={"skip": 0, "limit": 2},
    )

    assert resp.status_code == 200, resp.text
    assert len(resp.json()) == 2


@pytest.mark.asyncio
async def test_get_lotes_sin_permiso_devuelve_403(client, readonly_headers):
    resp = await client.get("/api/v1/previsionales/lotes", headers=readonly_headers)
    assert resp.status_code == 403


# ---------------------------------------------------------------------------
# GET /previsionales/lotes/{id}
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_get_lote_detalle(client, auditor_previsionales_headers, smlmv_2026):
    file_bytes = _construir_workbook_formalizacion([_FILA_VALIDA_1])
    post_resp = await client.post(
        "/api/v1/previsionales/lotes",
        headers=auditor_previsionales_headers,
        files={"file": ("LOTE_DETALLE.xlsx", file_bytes, _XLSX_CONTENT_TYPE)},
    )
    lote_id = post_resp.json()["id"]

    resp = await client.get(
        f"/api/v1/previsionales/lotes/{lote_id}", headers=auditor_previsionales_headers
    )

    assert resp.status_code == 200, resp.text
    assert resp.json()["id"] == lote_id
    assert resp.json()["nombre_archivo"] == "LOTE_DETALLE.xlsx"


@pytest.mark.asyncio
async def test_get_lote_detalle_404_para_id_inexistente(client, auditor_previsionales_headers):
    from uuid import uuid4

    resp = await client.get(
        f"/api/v1/previsionales/lotes/{uuid4()}", headers=auditor_previsionales_headers
    )
    assert resp.status_code == 404


# ---------------------------------------------------------------------------
# GET /previsionales/lotes/{id}/incapacidades  (filtros)
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_get_incapacidades_del_lote_filtro_errores(
    client, auditor_previsionales_headers, smlmv_2026
):
    file_bytes = _construir_workbook_formalizacion([_FILA_VALIDA_1, _FILA_MALFORMADA])
    post_resp = await client.post(
        "/api/v1/previsionales/lotes",
        headers=auditor_previsionales_headers,
        files={"file": ("LOTE_FILTROS.xlsx", file_bytes, _XLSX_CONTENT_TYPE)},
    )
    lote_id = post_resp.json()["id"]

    resp_todas = await client.get(
        f"/api/v1/previsionales/lotes/{lote_id}/incapacidades",
        headers=auditor_previsionales_headers,
    )
    assert resp_todas.status_code == 200, resp_todas.text
    assert len(resp_todas.json()) == 2

    resp_con_error = await client.get(
        f"/api/v1/previsionales/lotes/{lote_id}/incapacidades",
        headers=auditor_previsionales_headers,
        params={"errores": True},
    )
    assert resp_con_error.status_code == 200, resp_con_error.text
    con_error = resp_con_error.json()
    assert len(con_error) == 1
    assert con_error[0]["identificacion"] == "3000000003"
    assert con_error[0]["errores_carga"] is not None

    resp_sin_error = await client.get(
        f"/api/v1/previsionales/lotes/{lote_id}/incapacidades",
        headers=auditor_previsionales_headers,
        params={"errores": False},
    )
    assert resp_sin_error.status_code == 200, resp_sin_error.text
    sin_error = resp_sin_error.json()
    assert len(sin_error) == 1
    assert sin_error[0]["identificacion"] == "1000000001"


@pytest.mark.asyncio
async def test_get_incapacidades_lote_inexistente_devuelve_404(
    client, auditor_previsionales_headers
):
    from uuid import uuid4

    resp = await client.get(
        f"/api/v1/previsionales/lotes/{uuid4()}/incapacidades",
        headers=auditor_previsionales_headers,
    )
    assert resp.status_code == 404


# ---------------------------------------------------------------------------
# POST /previsionales/lotes/{id}/actualizar  ([GAP] re-cruce)
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_post_actualizar_lote_devuelve_501_no_implementado(
    client, auditor_previsionales_headers, smlmv_2026
):
    file_bytes = _construir_workbook_formalizacion([_FILA_VALIDA_1])
    post_resp = await client.post(
        "/api/v1/previsionales/lotes",
        headers=auditor_previsionales_headers,
        files={"file": ("LOTE_ACTUALIZAR.xlsx", file_bytes, _XLSX_CONTENT_TYPE)},
    )
    lote_id = post_resp.json()["id"]

    resp = await client.post(
        f"/api/v1/previsionales/lotes/{lote_id}/actualizar",
        headers=auditor_previsionales_headers,
    )

    assert resp.status_code == 501, resp.text
    assert "no implementado" in resp.json()["detail"].lower()


@pytest.mark.asyncio
async def test_post_actualizar_lote_inexistente_devuelve_404(
    client, auditor_previsionales_headers
):
    from uuid import uuid4

    resp = await client.post(
        f"/api/v1/previsionales/lotes/{uuid4()}/actualizar",
        headers=auditor_previsionales_headers,
    )
    assert resp.status_code == 404


@pytest.mark.asyncio
async def test_post_actualizar_lote_sin_permiso_devuelve_403(
    client, readonly_headers, auditor_previsionales_headers, smlmv_2026
):
    file_bytes = _construir_workbook_formalizacion([_FILA_VALIDA_1])
    post_resp = await client.post(
        "/api/v1/previsionales/lotes",
        headers=auditor_previsionales_headers,
        files={"file": ("LOTE_ACTUALIZAR_403.xlsx", file_bytes, _XLSX_CONTENT_TYPE)},
    )
    lote_id = post_resp.json()["id"]

    resp = await client.post(
        f"/api/v1/previsionales/lotes/{lote_id}/actualizar",
        headers=readonly_headers,
    )
    assert resp.status_code == 403


# ---------------------------------------------------------------------------
# GET /previsionales/lotes/{id}/arpis.xlsx
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_get_arpis_export_produce_xlsx_descargable_con_datos(
    client, auditor_previsionales_headers, smlmv_2026
):
    file_bytes = _construir_workbook_formalizacion([_FILA_VALIDA_1])
    post_resp = await client.post(
        "/api/v1/previsionales/lotes",
        headers=auditor_previsionales_headers,
        files={"file": ("LOTE_ARPIS.xlsx", file_bytes, _XLSX_CONTENT_TYPE)},
    )
    lote_id = post_resp.json()["id"]

    resp = await client.get(
        f"/api/v1/previsionales/lotes/{lote_id}/arpis.xlsx",
        headers=auditor_previsionales_headers,
    )

    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"] == _XLSX_CONTENT_TYPE
    assert "attachment" in resp.headers["content-disposition"]

    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    ws = wb["Cargue ARPIS"]

    header = [cell.value for cell in ws[3]]
    assert header[1] == "No. Identificación"

    fila_datos = [cell.value for cell in ws[4]]
    assert fila_datos[1] == "1000000001"  # No. Identificación
    assert fila_datos[6] == 1  # Tipo Ingreso -> INICIAL = codigo 1
    assert fila_datos[10] is not None  # Salario 1


@pytest.mark.asyncio
async def test_get_arpis_export_lote_inexistente_devuelve_404(
    client, auditor_previsionales_headers
):
    from uuid import uuid4

    resp = await client.get(
        f"/api/v1/previsionales/lotes/{uuid4()}/arpis.xlsx",
        headers=auditor_previsionales_headers,
    )
    assert resp.status_code == 404


@pytest.mark.asyncio
async def test_get_arpis_export_sin_permiso_devuelve_403(client, readonly_headers):
    from uuid import uuid4

    resp = await client.get(
        f"/api/v1/previsionales/lotes/{uuid4()}/arpis.xlsx",
        headers=readonly_headers,
    )
    assert resp.status_code == 403


# ---------------------------------------------------------------------------
# POST /previsionales/referencia/importar
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_post_referencia_importar_siniestros(client, auditor_previsionales_headers):
    file_bytes = _construir_workbook_siniestros(
        [["1000000001", "SIN-001", None, None, date(2026, 1, 1), None, "ABIERTO", None, None, None,
          None, None, None, None, None, None, None, None, "ARL", date(2026, 1, 2)]]
    )

    resp = await client.post(
        "/api/v1/previsionales/referencia/importar",
        headers=auditor_previsionales_headers,
        params={"tipo": "siniestros"},
        files={"file": ("SINIESTROS.xlsx", file_bytes, _XLSX_CONTENT_TYPE)},
    )

    assert resp.status_code == 200, resp.text
    data = resp.json()
    assert data["tipo"] == "siniestros"
    assert data["importados"] == 1


@pytest.mark.asyncio
async def test_post_referencia_importar_tipo_invalido_devuelve_422(
    client, auditor_previsionales_headers
):
    file_bytes = _construir_workbook_siniestros([])
    resp = await client.post(
        "/api/v1/previsionales/referencia/importar",
        headers=auditor_previsionales_headers,
        params={"tipo": "no-existe"},
        files={"file": ("SINIESTROS.xlsx", file_bytes, _XLSX_CONTENT_TYPE)},
    )
    assert resp.status_code == 422


@pytest.mark.asyncio
async def test_post_referencia_importar_sin_permiso_devuelve_403(client, readonly_headers):
    file_bytes = _construir_workbook_siniestros([])
    resp = await client.post(
        "/api/v1/previsionales/referencia/importar",
        headers=readonly_headers,
        params={"tipo": "siniestros"},
        files={"file": ("SINIESTROS.xlsx", file_bytes, _XLSX_CONTENT_TYPE)},
    )
    assert resp.status_code == 403


@pytest.mark.asyncio
async def test_post_referencia_importar_sin_token_devuelve_401(client):
    file_bytes = _construir_workbook_siniestros([])
    resp = await client.post(
        "/api/v1/previsionales/referencia/importar",
        params={"tipo": "siniestros"},
        files={"file": ("SINIESTROS.xlsx", file_bytes, _XLSX_CONTENT_TYPE)},
    )
    assert resp.status_code == 401


# ---------------------------------------------------------------------------
# GET /previsionales/parametros/smlmv
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_get_smlmv_lista_incluye_el_valor_sembrado(
    client, auditor_previsionales_headers, smlmv_2026
):
    resp = await client.get(
        "/api/v1/previsionales/parametros/smlmv", headers=auditor_previsionales_headers
    )

    assert resp.status_code == 200, resp.text
    data = resp.json()
    anos = {item["ano"] for item in data}
    assert 2026 in anos
    item_2026 = next(item for item in data if item["ano"] == 2026)
    assert Decimal(item_2026["valor"]) == Decimal("1750905.00")


@pytest.mark.asyncio
async def test_get_smlmv_sin_permiso_devuelve_403(client, readonly_headers):
    resp = await client.get("/api/v1/previsionales/parametros/smlmv", headers=readonly_headers)
    assert resp.status_code == 403


@pytest.mark.asyncio
async def test_get_smlmv_sin_token_devuelve_401(client):
    resp = await client.get("/api/v1/previsionales/parametros/smlmv")
    assert resp.status_code == 401


# ---------------------------------------------------------------------------
# ADMIN también debe poder acceder (permisos ya cableados en Task 0.1)
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_admin_puede_listar_lotes(client, admin_token_headers):
    resp = await client.get("/api/v1/previsionales/lotes", headers=admin_token_headers)
    assert resp.status_code == 200
