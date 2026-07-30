"""
Tests para los endpoints REST de liquidación y respuesta AFP previsional
(Task 4.3): `app/api/v1/endpoints/previsionales/liquidacion.py`.

Usa `client` (httpx AsyncClient contra la app real, `get_db` sobreescrito
con `db_session` -- ver tests/conftest.py), mismo patrón que
`test_endpoints_previsionales_lotes.py` (Task 4.1) para las fixtures de
usuario/headers, y mismo patrón que `test_respuesta_afp.py` (Task 3.4) para
sembrar un lote con archivo original archivado en `storage_backend` (vía
`metadata_["archivo_original_path"]`), necesario para GET /respuesta.xlsx.
"""
from __future__ import annotations

import io
from datetime import date
from decimal import Decimal
from uuid import uuid4

import openpyxl
import pytest
import pytest_asyncio
from openpyxl import Workbook

from app.core.security import create_access_token, get_password_hash
from app.core.storage_core import storage_backend
from app.models.previsionales.incapacidad_previsional import IncapacidadPrevisional
from app.models.previsionales.lote_previsional import LotePrevisional
from app.models.previsionales.periodo_previsional import PeriodoPrevisional
from app.models.previsionales.smlmv_parametros import SmlmvParametros
from app.models.usuario import Usuario
from app.utils.enums import EstadoUsuario, RolUsuario

_XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_ENCABEZADO_FORMALIZACION = [
    "Junio",
    "# IDENTIFICACION",
    "RADICADO",
    "FECHA RADICACION AFP",
    "TIPO INGRESO",
    "FECHA_INICIAL",
    "FECHA_FINAL",
    "No. DIAS",
]


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _construir_workbook_original() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "FORMALIZACION"
    ws.append(_ENCABEZADO_FORMALIZACION)
    ws.append([None, "1000000001", "R-001", None, "INICIAL", date(2026, 3, 1), date(2026, 3, 15), 15])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


async def _crear_usuario_con_rol(db_session, rol: RolUsuario, username: str) -> Usuario:
    usuario = Usuario(
        username=username,
        email=f"{username}@example.com",
        password_hash=get_password_hash("Test1234!"),
        nombre_completo=f"Usuario {rol.value}",
        rol=rol,
        estado=EstadoUsuario.ACTIVO,
    )
    db_session.add(usuario)
    await db_session.commit()
    await db_session.refresh(usuario)
    return usuario


def _headers_para(usuario: Usuario) -> dict:
    token = create_access_token(
        data={"sub": str(usuario.id), "token_version": usuario.token_version}
    )
    return {"Authorization": f"Bearer {token}"}


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------


@pytest_asyncio.fixture
async def smlmv_2026(db_session):
    smlmv = SmlmvParametros(ano=2026, valor=Decimal("1750905.00"), vigente_desde=date(2026, 1, 1))
    db_session.add(smlmv)
    await db_session.commit()
    return smlmv


@pytest_asyncio.fixture
async def auditor_previsionales_headers(db_session) -> dict:
    usuario = await _crear_usuario_con_rol(
        db_session, RolUsuario.AUDITOR_PREVISIONALES, "auditor_previsionales_liq_ep"
    )
    return _headers_para(usuario)


@pytest_asyncio.fixture
async def readonly_headers(db_session) -> dict:
    usuario = await _crear_usuario_con_rol(db_session, RolUsuario.READONLY, "readonly_previsionales_liq_ep")
    return _headers_para(usuario)


@pytest_asyncio.fixture
async def lote_con_periodos(db_session, smlmv_2026):
    """Un lote con una incapacidad no-duplicada con un periodo persistido
    (valores obsoletos a propósito, para verificar que POST /liquidar los
    recalcula de verdad), mismo patrón que
    `test_liquidacion_persistencia.py::lote_con_periodos`."""
    lote = LotePrevisional(nombre_archivo="RADICADOS_TEST_LIQUIDAR_EP.xlsx")
    db_session.add(lote)
    await db_session.flush()

    inc = IncapacidadPrevisional(
        lote_id=lote.id,
        identificacion="1000000001",
        fecha_inicial=date(2026, 3, 1),
        fecha_final=date(2026, 3, 15),
        valor_afp=Decimal("500000.00"),
        es_duplicado_interno=False,
    )
    db_session.add(inc)
    await db_session.flush()

    periodo = PeriodoPrevisional(
        incapacidad_id=inc.id,
        orden=1,
        fecha_inicio=date(2026, 3, 1),
        fecha_fin=date(2026, 3, 15),
        dias=15,
        ibc=Decimal("3000000.00"),
        salario=Decimal("3000000.00"),
        dias_cotizados=15,
        smlmv_aplicado=Decimal("1.00"),
        base_diaria=Decimal("1.00"),
        valor_segmento=Decimal("1.00"),
    )
    db_session.add(periodo)
    await db_session.commit()
    await db_session.refresh(inc)

    return {"lote": lote, "inc": inc}


@pytest_asyncio.fixture
async def lote_con_respuesta_pendiente(db_session):
    """Lote con el excel original archivado en storage (mismo camino que
    `lote_service.cargar_lote`) y una incapacidad avalada -- suficiente para
    ejercer GET /respuesta.xlsx end-to-end, mismo patrón que
    `test_respuesta_afp.py::lote_con_respuesta_pendiente`."""
    original_bytes = _construir_workbook_original()
    ruta_storage, _md5, _sha256, _size = storage_backend.upload_file(
        file_data=io.BytesIO(original_bytes),
        file_name="RADICADOS_TEST_RESPUESTA_EP.xlsx",
        content_type=_XLSX_CONTENT_TYPE,
        folder="lotes_previsionales",
    )

    lote = LotePrevisional(
        nombre_archivo="RADICADOS_TEST_RESPUESTA_EP.xlsx",
        metadata_={"archivo_original_path": ruta_storage},
    )
    db_session.add(lote)
    await db_session.flush()

    inc = IncapacidadPrevisional(
        lote_id=lote.id,
        identificacion="1000000001",
        radicado="R-001",
        fecha_inicial=date(2026, 3, 1),
        fecha_final=date(2026, 3, 15),
        aval="SI",
        es_duplicado_interno=False,
    )
    db_session.add(inc)
    await db_session.commit()
    await db_session.refresh(lote)

    return {"lote": lote, "inc": inc}


@pytest_asyncio.fixture
async def lote_sin_archivo_original(db_session):
    lote = LotePrevisional(nombre_archivo="SIN_ARCHIVO_EP.xlsx")
    db_session.add(lote)
    await db_session.commit()
    await db_session.refresh(lote)
    return lote


# ---------------------------------------------------------------------------
# POST /previsionales/lotes/{id}/liquidar
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_post_liquidar_recalcula_valor_auditado_y_retorna_conteo(
    client, auditor_previsionales_headers, lote_con_periodos
):
    lote = lote_con_periodos["lote"]
    inc = lote_con_periodos["inc"]

    resp = await client.post(
        f"/api/v1/previsionales/lotes/{lote.id}/liquidar",
        headers=auditor_previsionales_headers,
    )

    assert resp.status_code == 200, resp.text
    assert resp.json() == {"liquidadas": 1}

    # Verifica que valor_auditado realmente se actualizó en la BD (no solo
    # que el endpoint devolviera un número plausible).
    get_resp = await client.get(
        f"/api/v1/previsionales/lotes/{lote.id}/incapacidades",
        headers=auditor_previsionales_headers,
    )
    assert get_resp.status_code == 200, get_resp.text
    incapacidades = get_resp.json()
    actualizada = next(i for i in incapacidades if i["id"] == str(inc.id))
    assert actualizada["valor_auditado"] is not None
    assert Decimal(actualizada["valor_auditado"]) != Decimal("1.00")


@pytest.mark.asyncio
async def test_post_liquidar_lote_inexistente_devuelve_404(client, auditor_previsionales_headers):
    resp = await client.post(
        f"/api/v1/previsionales/lotes/{uuid4()}/liquidar",
        headers=auditor_previsionales_headers,
    )
    assert resp.status_code == 404


@pytest.mark.asyncio
async def test_post_liquidar_sin_permiso_devuelve_403(client, readonly_headers, lote_con_periodos):
    lote = lote_con_periodos["lote"]

    resp = await client.post(
        f"/api/v1/previsionales/lotes/{lote.id}/liquidar",
        headers=readonly_headers,
    )
    assert resp.status_code == 403


@pytest.mark.asyncio
async def test_post_liquidar_sin_token_devuelve_401(client, lote_con_periodos):
    lote = lote_con_periodos["lote"]

    resp = await client.post(f"/api/v1/previsionales/lotes/{lote.id}/liquidar")
    assert resp.status_code == 401


# ---------------------------------------------------------------------------
# GET /previsionales/lotes/{id}/respuesta.xlsx
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_get_respuesta_produce_xlsx_descargable_con_columnas_aval_observacion(
    client, auditor_previsionales_headers, lote_con_respuesta_pendiente
):
    lote = lote_con_respuesta_pendiente["lote"]

    resp = await client.get(
        f"/api/v1/previsionales/lotes/{lote.id}/respuesta.xlsx",
        headers=auditor_previsionales_headers,
    )

    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"] == _XLSX_CONTENT_TYPE
    assert "attachment" in resp.headers["content-disposition"]

    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    ws = wb["FORMALIZACION"]
    encabezados = [c.value for c in ws[1]]
    assert encabezados[-2:] == ["AVAL", "OBSERVACION"]

    fila = [c.value for c in ws[2]]
    assert fila[1] == "1000000001"
    assert fila[-2] == "SI"  # AVAL


@pytest.mark.asyncio
async def test_get_respuesta_lote_inexistente_devuelve_404(client, auditor_previsionales_headers):
    resp = await client.get(
        f"/api/v1/previsionales/lotes/{uuid4()}/respuesta.xlsx",
        headers=auditor_previsionales_headers,
    )
    assert resp.status_code == 404


@pytest.mark.asyncio
async def test_get_respuesta_lote_sin_archivo_original_devuelve_400(
    client, auditor_previsionales_headers, lote_sin_archivo_original
):
    resp = await client.get(
        f"/api/v1/previsionales/lotes/{lote_sin_archivo_original.id}/respuesta.xlsx",
        headers=auditor_previsionales_headers,
    )
    assert resp.status_code == 400


@pytest.mark.asyncio
async def test_get_respuesta_sin_permiso_devuelve_403(client, readonly_headers, lote_con_respuesta_pendiente):
    lote = lote_con_respuesta_pendiente["lote"]

    resp = await client.get(
        f"/api/v1/previsionales/lotes/{lote.id}/respuesta.xlsx",
        headers=readonly_headers,
    )
    assert resp.status_code == 403


@pytest.mark.asyncio
async def test_get_respuesta_sin_token_devuelve_401(client, lote_con_respuesta_pendiente):
    lote = lote_con_respuesta_pendiente["lote"]

    resp = await client.get(f"/api/v1/previsionales/lotes/{lote.id}/respuesta.xlsx")
    assert resp.status_code == 401
