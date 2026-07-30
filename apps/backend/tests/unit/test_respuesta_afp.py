"""
Tests para `respuesta_afp.py` (Task 3.4, segunda mitad) -- generador del
excel de respuesta al fondo (AVAL + OBSERVACION agregadas al libro original).

Usa `db_session` (Postgres async real) y `storage_backend` REAL (filesystem,
igual que `test_lote_service.py` -- no se mockea, el contenedor de test ya
corre con `STORAGE_BACKEND=filesystem`): se construye un workbook sintético
en memoria con la forma mínima de la hoja FORMALIZACION, se sube vía
`storage_backend.upload_file` (mismo camino que usa `lote_service.cargar_lote`
para archivar el original), y se referencia desde
`lote.metadata_["archivo_original_path"]`.
"""
import io
from datetime import date

import openpyxl
import pytest
import pytest_asyncio
from openpyxl import Workbook

from app.core.exceptions import BadRequestException, NotFoundException
from app.core.storage_core import storage_backend
from app.models.previsionales.incapacidad_previsional import IncapacidadPrevisional
from app.models.previsionales.lote_previsional import LotePrevisional
from app.models.previsionales.senal_auditoria_previsional import SenalAuditoriaPrevisional
from app.services.previsionales.respuesta_afp import generar_respuesta

_ENCABEZADO_FORMALIZACION = [
    "Junio",
    "# IDENTIFICACION",
    "RADICADO",
    "FECHA RADICACION AFP",
    "TIPO INGRESO",
    "FECHA_INICIAL",
    "FECHA_FINAL",
    "No. DIAS",
]

# (identificacion, radicado, fecha_inicial, fecha_final) -- suficiente para
# el match por (identificacion, fecha_inicial); las demas columnas son
# irrelevantes para este test pero se dejan para simular un libro real.
_FILAS = [
    ("1000000001", "R-001", date(2026, 3, 1), date(2026, 3, 15)),  # avalada SI
    ("1000000002", "R-002", date(2026, 3, 1), date(2026, 3, 10)),  # avalada NO
    ("1000000003", "R-003", date(2026, 3, 1), date(2026, 3, 5)),  # duplicado interno
    ("1000000004", "R-004", date(2026, 3, 1), date(2026, 3, 8)),  # sin auditar (aval None)
]


def _construir_workbook_original() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "FORMALIZACION"
    ws.append(_ENCABEZADO_FORMALIZACION)
    for identificacion, radicado, f_ini, f_fin in _FILAS:
        ws.append([None, identificacion, radicado, None, "INICIAL", f_ini, f_fin, None])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


@pytest_asyncio.fixture
async def lote_con_respuesta_pendiente(db_session):
    original_bytes = _construir_workbook_original()
    ruta_storage, _md5, _sha256, _size = storage_backend.upload_file(
        file_data=io.BytesIO(original_bytes),
        file_name="RADICADOS_TEST_RESPUESTA.xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        folder="lotes_previsionales",
    )

    lote = LotePrevisional(
        nombre_archivo="RADICADOS_TEST_RESPUESTA.xlsx",
        metadata_={"archivo_original_path": ruta_storage},
    )
    db_session.add(lote)
    await db_session.flush()

    inc_si = IncapacidadPrevisional(
        lote_id=lote.id,
        identificacion="1000000001",
        radicado="R-001",
        fecha_inicial=date(2026, 3, 1),
        fecha_final=date(2026, 3, 15),
        aval="SI",
        dia_181_alfa=date(2026, 8, 1),
        es_duplicado_interno=False,
    )
    inc_no = IncapacidadPrevisional(
        lote_id=lote.id,
        identificacion="1000000002",
        radicado="R-002",
        fecha_inicial=date(2026, 3, 1),
        fecha_final=date(2026, 3, 10),
        aval="NO",
        motivo_no_aval="Diagnostico no compatible con incapacidad prolongada",
        es_duplicado_interno=False,
    )
    inc_duplicada = IncapacidadPrevisional(
        lote_id=lote.id,
        identificacion="1000000003",
        radicado="R-003",
        fecha_inicial=date(2026, 3, 1),
        fecha_final=date(2026, 3, 5),
        aval="SI",  # aunque tenga aval, es_duplicado_interno la excluye igual
        es_duplicado_interno=True,
    )
    inc_sin_auditar = IncapacidadPrevisional(
        lote_id=lote.id,
        identificacion="1000000004",
        radicado="R-004",
        fecha_inicial=date(2026, 3, 1),
        fecha_final=date(2026, 3, 8),
        aval=None,
        es_duplicado_interno=False,
    )
    db_session.add_all([inc_si, inc_no, inc_duplicada, inc_sin_auditar])
    await db_session.flush()

    senal_ab = SenalAuditoriaPrevisional(
        incapacidad_previsional_id=inc_si.id,
        codigo="AB",
        nombre="Dia 181 (auditado)",
        estado="OK",
        valor="DIA 181 2026-08-01",
        detalle=None,
    )
    db_session.add(senal_ab)
    await db_session.commit()

    return {
        "lote": lote,
        "inc_si": inc_si,
        "inc_no": inc_no,
        "inc_duplicada": inc_duplicada,
        "inc_sin_auditar": inc_sin_auditar,
    }


def _leer_fila_por_identificacion(ws, identificacion: str) -> dict:
    encabezados = [c.value for c in ws[1]]
    idx_ident = encabezados.index("# IDENTIFICACION") + 1
    idx_aval = len(encabezados) - 1  # AVAL fue agregada como penúltima columna
    idx_obs = len(encabezados)  # OBSERVACION es la última
    for row in ws.iter_rows(min_row=2):
        if row[idx_ident - 1].value == identificacion:
            return {
                "aval": row[idx_aval - 1].value,
                "observacion": row[idx_obs - 1].value,
                "radicado": row[2].value,  # columna original C, sin modificar
            }
    raise AssertionError(f"Fila con identificacion={identificacion!r} no encontrada")


@pytest.mark.asyncio
async def test_agrega_columnas_aval_y_observacion_al_encabezado(db_session, lote_con_respuesta_pendiente):
    lote = lote_con_respuesta_pendiente["lote"]

    resultado_bytes = await generar_respuesta(db_session, lote.id)

    wb = openpyxl.load_workbook(io.BytesIO(resultado_bytes))
    ws = wb["FORMALIZACION"]
    encabezados = [c.value for c in ws[1]]
    assert encabezados[-2:] == ["AVAL", "OBSERVACION"]
    # Las columnas originales no se tocaron.
    assert encabezados[:8] == _ENCABEZADO_FORMALIZACION


@pytest.mark.asyncio
async def test_fila_avalada_si_trae_aval_y_observacion_desde_senal_ab(
    db_session, lote_con_respuesta_pendiente
):
    lote = lote_con_respuesta_pendiente["lote"]

    resultado_bytes = await generar_respuesta(db_session, lote.id)

    wb = openpyxl.load_workbook(io.BytesIO(resultado_bytes))
    ws = wb["FORMALIZACION"]
    fila = _leer_fila_por_identificacion(ws, "1000000001")
    assert fila["aval"] == "SI"
    assert fila["observacion"] == "DIA 181 2026-08-01"
    assert fila["radicado"] == "R-001"  # dato original intacto


@pytest.mark.asyncio
async def test_fila_avalada_no_trae_aval_y_motivo_no_aval(db_session, lote_con_respuesta_pendiente):
    lote = lote_con_respuesta_pendiente["lote"]

    resultado_bytes = await generar_respuesta(db_session, lote.id)

    wb = openpyxl.load_workbook(io.BytesIO(resultado_bytes))
    ws = wb["FORMALIZACION"]
    fila = _leer_fila_por_identificacion(ws, "1000000002")
    assert fila["aval"] == "NO"
    assert fila["observacion"] == "Diagnostico no compatible con incapacidad prolongada"


@pytest.mark.asyncio
async def test_fila_duplicado_interno_queda_sin_aval_ni_observacion(
    db_session, lote_con_respuesta_pendiente
):
    lote = lote_con_respuesta_pendiente["lote"]

    resultado_bytes = await generar_respuesta(db_session, lote.id)

    wb = openpyxl.load_workbook(io.BytesIO(resultado_bytes))
    ws = wb["FORMALIZACION"]
    fila = _leer_fila_por_identificacion(ws, "1000000003")
    assert fila["aval"] is None
    assert fila["observacion"] is None
    # La fila SIGUE presente en el libro (no se borra), solo sin las 2 celdas.
    assert fila["radicado"] == "R-003"


@pytest.mark.asyncio
async def test_fila_sin_auditar_aval_none_queda_en_blanco(db_session, lote_con_respuesta_pendiente):
    lote = lote_con_respuesta_pendiente["lote"]

    resultado_bytes = await generar_respuesta(db_session, lote.id)

    wb = openpyxl.load_workbook(io.BytesIO(resultado_bytes))
    ws = wb["FORMALIZACION"]
    fila = _leer_fila_por_identificacion(ws, "1000000004")
    assert fila["aval"] is None
    assert fila["observacion"] is None


@pytest.mark.asyncio
async def test_no_reconstruye_el_libro_todas_las_filas_originales_se_conservan(
    db_session, lote_con_respuesta_pendiente
):
    lote = lote_con_respuesta_pendiente["lote"]

    resultado_bytes = await generar_respuesta(db_session, lote.id)

    wb = openpyxl.load_workbook(io.BytesIO(resultado_bytes))
    ws = wb["FORMALIZACION"]
    # Encabezado (1) + 4 filas de datos, ninguna se borró ni se agregó.
    assert ws.max_row == 5
    identificaciones = [ws.cell(row=r, column=2).value for r in range(2, ws.max_row + 1)]
    assert identificaciones == ["1000000001", "1000000002", "1000000003", "1000000004"]


@pytest.mark.asyncio
async def test_fila_avalada_si_con_senal_ab_pendiente_no_filtra_texto_interno(
    db_session, lote_con_respuesta_pendiente
):
    """Fix-round Finding 1: si aval="SI" pero la señal AB persistida sigue
    PENDIENTE (valor=None, dia_181_alfa aún no calculado por el auditor),
    la OBSERVACION debe quedar en blanco -- NUNCA el texto interno de
    `detalle` ("Aun no hay dia_181_alfa calculado por el auditor")."""
    lote = lote_con_respuesta_pendiente["lote"]
    inc_si = lote_con_respuesta_pendiente["inc_si"]

    # Reemplaza la señal AB "OK" del fixture por una PENDIENTE (valor=None,
    # detalle=mensaje interno de depuración) -- exactamente lo que produce
    # `auditoria_rules.regla_ab_observacion` mientras `dia_181_alfa` no ha
    # sido fijado por el auditor (el estado por defecto de casi toda fila
    # hoy en el sistema).
    from sqlalchemy import delete

    await db_session.execute(
        delete(SenalAuditoriaPrevisional).where(
            SenalAuditoriaPrevisional.incapacidad_previsional_id == inc_si.id
        )
    )
    senal_pendiente = SenalAuditoriaPrevisional(
        incapacidad_previsional_id=inc_si.id,
        codigo="AB",
        nombre="Dia 181 (auditado)",
        estado="PENDIENTE",
        valor=None,
        detalle="Aun no hay dia_181_alfa calculado por el auditor",
    )
    db_session.add(senal_pendiente)
    await db_session.commit()

    resultado_bytes = await generar_respuesta(db_session, lote.id)

    wb = openpyxl.load_workbook(io.BytesIO(resultado_bytes))
    ws = wb["FORMALIZACION"]
    fila = _leer_fila_por_identificacion(ws, "1000000001")
    assert fila["aval"] == "SI"
    assert fila["observacion"] is None
    assert fila["observacion"] != "Aun no hay dia_181_alfa calculado por el auditor"


@pytest.mark.asyncio
async def test_multiples_filas_avaladas_si_reciben_su_propia_observacion_batch(
    db_session, lote_con_respuesta_pendiente
):
    """Fix-round Finding 2: la observación de cada fila avalada-SI se
    resuelve con un lookup batch (una sola query para todo el lote), no una
    query por fila. Prueba de comportamiento: 2 incapacidades avaladas-SI en
    el mismo lote, cada una con su propia señal AB "OK", deben terminar con
    SU PROPIO texto -- no el de la otra, ni un texto compartido/mezclado."""
    lote = lote_con_respuesta_pendiente["lote"]
    inc_no = lote_con_respuesta_pendiente["inc_no"]

    # Convierte la fila "avalada NO" del fixture (identificacion
    # 1000000002) en una segunda fila avalada-SI, con su propia señal AB
    # distinta de la de inc_si (identificacion 1000000001, DIA 181
    # 2026-08-01 según el fixture).
    inc_no.aval = "SI"
    inc_no.motivo_no_aval = None
    db_session.add(inc_no)
    await db_session.flush()

    senal_ab_segunda = SenalAuditoriaPrevisional(
        incapacidad_previsional_id=inc_no.id,
        codigo="AB",
        nombre="Dia 181 (auditado)",
        estado="OK",
        valor="DIA 181 2026-09-15",
        detalle=None,
    )
    db_session.add(senal_ab_segunda)
    await db_session.commit()

    resultado_bytes = await generar_respuesta(db_session, lote.id)

    wb = openpyxl.load_workbook(io.BytesIO(resultado_bytes))
    ws = wb["FORMALIZACION"]

    fila_1 = _leer_fila_por_identificacion(ws, "1000000001")
    fila_2 = _leer_fila_por_identificacion(ws, "1000000002")
    assert fila_1["aval"] == "SI"
    assert fila_1["observacion"] == "DIA 181 2026-08-01"
    assert fila_2["aval"] == "SI"
    assert fila_2["observacion"] == "DIA 181 2026-09-15"
    assert fila_1["observacion"] != fila_2["observacion"]


@pytest.mark.asyncio
async def test_lote_inexistente_lanza_not_found(db_session):
    from uuid import uuid4

    with pytest.raises(NotFoundException):
        await generar_respuesta(db_session, uuid4())


@pytest.mark.asyncio
async def test_lote_sin_archivo_original_lanza_bad_request(db_session):
    lote = LotePrevisional(nombre_archivo="SIN_ARCHIVO.xlsx")
    db_session.add(lote)
    await db_session.commit()
    await db_session.refresh(lote)

    with pytest.raises(BadRequestException):
        await generar_respuesta(db_session, lote.id)
