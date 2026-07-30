"""
Tests para `lote_excel_reader.py` (Task 3.1).

Usa el fixture cifrado real de Task 0.3
(tests/fixtures/afp_formalizacion_encrypted.xlsx, password "test1234") para
probar contraseña correcta/incorrecta/ausente, y libros sintéticos en
memoria para el resto de los casos (plano, sin hoja FORMALIZACION, no es un
Excel válido).
"""
import io
from pathlib import Path

import openpyxl
import pytest

from app.core.exceptions import BadRequestException
from app.services.previsionales.lote_excel_reader import (
    cargar_hoja_desde_bytes_planos,
    cargar_hoja_formalizacion,
    descifrar_si_aplica,
)

FIXTURE_PATH = Path(__file__).parent.parent / "fixtures" / "afp_formalizacion_encrypted.xlsx"
FIXTURE_PASSWORD = "test1234"


def _libro_plano_con_formalizacion() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "FORMALIZACION"
    ws["A1"] = "encabezado"
    ws["A2"] = "dato"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _libro_plano_sin_formalizacion() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "OTRA_HOJA"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Contraseña correcta / incorrecta / ausente (fixture cifrado real)
# ---------------------------------------------------------------------------


def test_password_correcta_desencripta_y_lee_formalizacion():
    file_bytes = FIXTURE_PATH.read_bytes()
    ws = cargar_hoja_formalizacion(file_bytes, FIXTURE_PASSWORD)

    assert ws.title == "FORMALIZACION"
    header = [cell.value for cell in ws[1]]
    assert header == ["TIPO_DOC", "NUMERO_DOC", "NOMBRE_AFILIADO", "MONTO"]


def test_password_incorrecta_lanza_bad_request():
    file_bytes = FIXTURE_PATH.read_bytes()
    with pytest.raises(BadRequestException) as exc_info:
        cargar_hoja_formalizacion(file_bytes, "password-equivocada")

    assert "contraseña incorrecta" in str(exc_info.value)


def test_sin_password_en_archivo_cifrado_lanza_bad_request():
    file_bytes = FIXTURE_PATH.read_bytes()
    with pytest.raises(BadRequestException) as exc_info:
        cargar_hoja_formalizacion(file_bytes, None)

    assert "requiere contraseña" in str(exc_info.value)


def test_password_incorrecta_nunca_aparece_en_mensaje_de_error():
    file_bytes = FIXTURE_PATH.read_bytes()
    password_secreta = "super-secreta-nunca-debe-salir"
    with pytest.raises(BadRequestException) as exc_info:
        cargar_hoja_formalizacion(file_bytes, password_secreta)

    assert password_secreta not in str(exc_info.value)
    assert password_secreta not in repr(exc_info.value)


# ---------------------------------------------------------------------------
# Archivo plano (no cifrado) pasa sin tocar msoffcrypto
# ---------------------------------------------------------------------------


def test_archivo_plano_no_cifrado_pasa_sin_password():
    file_bytes = _libro_plano_con_formalizacion()
    ws = cargar_hoja_formalizacion(file_bytes, None)

    assert ws.title == "FORMALIZACION"
    assert ws["A1"].value == "encabezado"


def test_archivo_plano_no_cifrado_ignora_password_dado():
    """Un .xlsx plano nunca pasa por msoffcrypto, incluso si se da password."""
    file_bytes = _libro_plano_con_formalizacion()
    ws = cargar_hoja_formalizacion(file_bytes, "password-que-no-deberia-importar")

    assert ws.title == "FORMALIZACION"


def test_descifrar_si_aplica_es_passthrough_para_archivo_plano():
    file_bytes = _libro_plano_con_formalizacion()
    assert descifrar_si_aplica(file_bytes, None) == file_bytes


# ---------------------------------------------------------------------------
# Hoja FORMALIZACION ausente / archivo inválido
# ---------------------------------------------------------------------------


def test_libro_sin_hoja_formalizacion_lanza_bad_request():
    file_bytes = _libro_plano_sin_formalizacion()
    with pytest.raises(BadRequestException) as exc_info:
        cargar_hoja_formalizacion(file_bytes, None)

    assert "FORMALIZACION" in str(exc_info.value)


def test_archivo_no_excel_lanza_bad_request_con_mensaje_especifico():
    """Ni OLE2 ni un zip OOXML válido -> mensaje distinto a los de contraseña."""
    file_bytes = b"esto no es ni un OLE2 ni un zip valido" * 5
    with pytest.raises(BadRequestException) as exc_info:
        cargar_hoja_formalizacion(file_bytes, None)

    assert "no es un Excel válido" in str(exc_info.value)


def test_cargar_hoja_desde_bytes_planos_reusa_bytes_ya_desencriptados():
    """Permite al servicio reusar los bytes decifrados sin re-desencriptar."""
    file_bytes = FIXTURE_PATH.read_bytes()
    from app.services.previsionales.lote_excel_reader import descifrar_si_aplica as _d

    bytes_planos = _d(file_bytes, FIXTURE_PASSWORD)
    ws = cargar_hoja_desde_bytes_planos(bytes_planos)
    assert ws.title == "FORMALIZACION"
