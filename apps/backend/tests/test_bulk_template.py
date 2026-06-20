"""Tests for Excel template generation endpoint (radicar masiva)."""
import io
import pytest
from openpyxl import load_workbook
from httpx import AsyncClient

EXPECTED_HEADERS = [
    "numero_documento", "tipo_documento", "empleado_nombres", "empleado_apellidos",
    "tipo_enfermedad", "fecha_inicio", "fecha_fin", "dias_totales", "diagnostico_cie10",
    "descripcion_diagnostico", "nombre_medico", "registro_medico", "ips", "valor_dia",
    "prorroga", "observaciones",
]


@pytest.mark.asyncio
async def test_template_headers_only(client: AsyncClient, empresa_user_token):
    resp = await client.post(
        "/api/v1/incapacidades/radicar-masiva/plantilla",
        json={"empleado_ids": []},
        headers={"Authorization": f"Bearer {empresa_user_token}"},
    )
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    headers = [c.value for c in ws[1]]
    assert headers == EXPECTED_HEADERS
    assert ws.max_row == 1


@pytest.mark.asyncio
async def test_template_prefilled(client: AsyncClient, empresa_user_token, test_empleado):
    resp = await client.post(
        "/api/v1/incapacidades/radicar-masiva/plantilla",
        json={"empleado_ids": [str(test_empleado.id)]},
        headers={"Authorization": f"Bearer {empresa_user_token}"},
    )
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    assert ws.max_row == 2
    assert ws.cell(row=2, column=1).value == test_empleado.numero_documento
