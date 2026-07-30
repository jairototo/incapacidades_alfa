"""
Adaptador de datos de referencia previsionales (Task 3.2).

Importa desde Excel los tres tipos de datos de referencia que el motor de
auditoría (Task 1.4, `auditoria_rules.py`) usa como fuente de solo lectura
al construir el `ContextoAuditoria` de un lote (Task 3.3):

- `SiniestroPrevisional` (hoja SINIESTROS): siniestros laborales ya
  conocidos por el asegurador.
- `SolicitudPrevisional` (hoja SOLICITUDES): incapacidades ya radicadas en
  el sistema externo ARPIS/AFP.
- `IteHistorico` (hoja "LISTADO ITE DIA"): histórico de pagos de ITE, usado
  para el anti-doble-pago (regla AT).

La vista SQL real del asegurador todavía no existe (verificado esta sesión
directamente contra el workbook real `RADICADOS_AUDITORIA_20260706.xlsx`,
que NO está cifrado — abre como zip/OOXML plano, magic bytes `PK\\x03\\x04`,
sin contraseña). Este adaptador es un puente temporal detrás de
`FuenteReferenciaPrevisional` (Protocol) — mismo seam que
`app/services/integracion/servialfa_client.py` / `sicat_client.py`. El día
que la vista SQL exista, se agrega una segunda implementación del Protocol
(p.ej. `SqlReferenciaAdapter`) y el código llamador (Task 3.3) no cambia —
solo el adaptador que se inyecta.

## `origen_dato`: GAP conocido

Ninguno de los tres modelos (Task 2.1: `SiniestroPrevisional`,
`SolicitudPrevisional`, `IteHistorico`) tiene una columna `origen_dato` (en
el sentido `EXCEL` | `VISTA_SQL`) para distinguir el origen de una fila
importada. Se deja como GAP explícito — no se agrega una migración aquí,
fuera de alcance de esta tarea — documentado en el reporte de Task 3.2.
Mientras tanto, no hay forma de distinguir en BD si una fila vino de este
adaptador Excel o (en el futuro) de la vista SQL real.

## Idempotencia

Ninguna de las 3 tablas tiene una restricción UNIQUE en BD sobre su llave
natural, así que una re-importación NO fallaría con `IntegrityError` — pero
insertar ciegamente duplicaría filas en cada reimport. Este adaptador
mitiga esto a nivel de aplicación, reutilizando los métodos de carga masiva
ya existentes de Task 2.2 (mismo principio "nunca N+1" que CLAUDE.md exige
tanto en lectura como en escritura):

- Siniestros: llave natural `(identificacion, numero_siniestro)`. Se
  precargan los siniestros existentes de las identificaciones del archivo
  con `get_by_identificaciones` y se omiten filas cuya llave ya existe.
- Solicitudes: llave natural `(identificacion, radicado)`. Misma
  estrategia con `get_by_identificaciones`. Si `radicado` es `None`
  (columna B vacía) no hay llave natural confiable — esas filas SIEMPRE
  se insertan (no se puede deduplicar sin un radicado).
- ITE histórico: llave natural `(identificacion, fecha_inicial)` — exacta
  para la que ya existe `get_existentes` (bulk exists-check de Task 2.2,
  pensada precisamente para esto).

Esto hace que re-importar el mismo archivo sea un no-op para las filas ya
importadas (no crecen sin límite en cada corrida), sin depender de una
restricción UNIQUE que no existe en el esquema. También se deduplica
DENTRO del mismo archivo (misma llave repetida dos veces en la misma
hoja) con el mismo criterio.

## Errores por fila

Un error de parseo en una fila (celda inválida, capturado como
`CeldaInvalidaError`, o campo requerido vacío) NO aborta el resto de la
hoja — se omite y se sigue, mismo principio que Task 3.1
(`lote_service.cargar_lote`). Se loguea cada fila omitida (`logger.warning`)
para trazabilidad operativa, pero el contrato del Protocol
(`-> int`, cantidad importada) no expone el detalle de errores — no existe
hoy un modelo ni un consumidor downstream obvio para persistirlos. A
diferencia de `cargar_lote` (que usa un SAVEPOINT por fila para aislar
también errores de BD en el flush), aquí el flush es ÚNICO por hoja
(`bulk_create_flushed`): un fallo de BD a nivel de flush (p.ej. overflow de
columna) abortaría toda la hoja, no solo la fila — ver el reporte de esta
tarea para el razonamiento de este trade-off.
"""
import io
from typing import Any, Protocol
from zipfile import BadZipFile

import openpyxl
from loguru import logger
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.exceptions import BadRequestException
from app.db.repositories.previsionales.ite_historico_repository import (
    ite_historico_repository,
)
from app.db.repositories.previsionales.siniestro_previsional_repository import (
    siniestro_previsional_repository,
)
from app.db.repositories.previsionales.solicitud_previsional_repository import (
    solicitud_previsional_repository,
)
from app.services.previsionales.excel_common import (
    CeldaInvalidaError,
    es_vacio,
    parse_fecha,
    texto,
    validar_encabezados,
)

_HOJA_SINIESTROS = "SINIESTROS"
_HOJA_SOLICITUDES = "SOLICITUDES"
_HOJA_ITE = "LISTADO ITE DIA"
_FILA_ENCABEZADO = 1

# Columnas críticas (mínimo exigido por el brief) más las que este
# adaptador efectivamente lee — validar TODO lo que se lee, no solo el
# mínimo, para no confiar nunca en un layout de columnas no verificado.
_ENCABEZADOS_SINIESTROS = {
    "A": "Número Id",
    "B": "No. Siniestro",
    "E": "Fecha Siniestro",
    "G": "Estado siniestro",
    "S": "Origen siniestro",
    "T": "Fecha Aviso",
}
_ENCABEZADOS_SOLICITUDES = {
    "B": "No. Solicitud Auditoria",
    "D": "Número Id",
    "N": "Dia 181",
    "R": "Observaciones",
    "AF": "FECHA_CRIE",
}
_ENCABEZADOS_ITE = {
    "F": "Número identificacion",
    "I": "Fecha inicial",
    "J": "Fecha final",
}


class FuenteReferenciaPrevisional(Protocol):
    """
    Seam de integración para los datos de referencia previsionales.

    Cualquier implementación (Excel hoy, vista SQL real el día que exista)
    expone estos tres métodos — el llamador (Task 3.3,
    `construir_contexto`) nunca sabe cuál está detrás, solo cambia qué
    instancia se inyecta.
    """

    async def importar_siniestros(self, db: AsyncSession, file_bytes: bytes) -> int:
        """Importa la hoja SINIESTROS. Retorna la cantidad de filas nuevas persistidas."""
        ...

    async def importar_solicitudes(self, db: AsyncSession, file_bytes: bytes) -> int:
        """Importa la hoja SOLICITUDES. Retorna la cantidad de filas nuevas persistidas."""
        ...

    async def importar_ite_historico(self, db: AsyncSession, file_bytes: bytes) -> int:
        """Importa la hoja LISTADO ITE DIA. Retorna la cantidad de filas nuevas persistidas."""
        ...


def _get(fila: tuple, idx: int) -> Any:
    """Acceso defensivo a `fila[idx]`: `None` si la fila es más corta que `idx`."""
    return fila[idx] if idx < len(fila) else None


def _cargar_hoja(file_bytes: bytes, nombre_hoja: str) -> Worksheet:
    """
    Abre el libro (se asume plano, no cifrado — confirmado contra el
    workbook real de esta sesión) y retorna la hoja pedida.

    Raises:
        BadRequestException: "no es un Excel válido" si `file_bytes` no es
            un .xlsx válido; "Hoja ... no encontrada" si el libro no trae
            esa hoja.
    """
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    except BadZipFile:
        raise BadRequestException("no es un Excel válido")
    except Exception:
        raise BadRequestException("no es un Excel válido")

    if nombre_hoja not in wb.sheetnames:
        raise BadRequestException(f"Hoja {nombre_hoja!r} no encontrada")

    return wb[nombre_hoja]


class ExcelReferenciaAdapter:
    """
    Implementación de `FuenteReferenciaPrevisional` respaldada por las
    hojas de referencia del Excel del asegurador (SINIESTROS, SOLICITUDES,
    "LISTADO ITE DIA" de `RADICADOS_AUDITORIA_*.xlsx`) — puente temporal
    mientras no exista la vista SQL real.
    """

    # -- SINIESTROS -----------------------------------------------------

    @staticmethod
    def _parsear_fila_siniestro(fila: tuple, num_fila: int) -> dict[str, Any] | None:
        """
        Mapeo SINIESTROS: A=identificacion, B=numero_siniestro,
        E=fecha_siniestro, G=estado, S=origen, T=fecha_aviso. El resto de
        columnas (C, D, F, H..R, U) son informativas o no tienen columna
        equivalente en `SiniestroPrevisional` — no se persisten (ver
        reporte de Task 3.2 para el detalle columna por columna). La
        columna U (`FECHA SINIESTRO < FECHA AVISO`) es un booleano
        precalculado en la fuente que se IGNORA deliberadamente: no hay
        columna en el modelo para guardarlo, y de necesitarse se recalcula
        trivialmente desde E/T en la capa de reglas, nunca se confía en el
        valor de la fuente.

        Retorna `None` (fila omitida, sin contar como error) si faltan los
        campos requeridos (`identificacion`/`numero_siniestro` vacíos).
        Propaga `CeldaInvalidaError` si una fecha no se puede parsear.
        """
        identificacion = texto(_get(fila, 0))  # A
        numero_siniestro = texto(_get(fila, 1))  # B
        fecha_siniestro = parse_fecha(_get(fila, 4))  # E
        estado = texto(_get(fila, 6))  # G
        origen = texto(_get(fila, 18))  # S
        fecha_aviso = parse_fecha(_get(fila, 19))  # T

        if es_vacio(identificacion) or es_vacio(numero_siniestro):
            logger.warning(
                f"SINIESTROS fila {num_fila}: identificación o número de "
                "siniestro vacío — fila omitida (requeridos por "
                "SiniestroPrevisional.identificacion/numero_siniestro)."
            )
            return None

        return {
            "identificacion": identificacion,
            "numero_siniestro": numero_siniestro,
            "origen": origen,
            "estado": estado,
            "fecha_aviso": fecha_aviso,
            "fecha_siniestro": fecha_siniestro,
        }

    async def importar_siniestros(self, db: AsyncSession, file_bytes: bytes) -> int:
        ws = _cargar_hoja(file_bytes, _HOJA_SINIESTROS)
        validar_encabezados(ws, _ENCABEZADOS_SINIESTROS, fila=_FILA_ENCABEZADO)

        candidatas: list[dict[str, Any]] = []
        for i, fila in enumerate(
            ws.iter_rows(min_row=_FILA_ENCABEZADO + 1, values_only=True),
            start=_FILA_ENCABEZADO + 1,
        ):
            try:
                campos = self._parsear_fila_siniestro(fila, i)
            except CeldaInvalidaError as exc:
                logger.warning(f"SINIESTROS fila {i}: error de parseo, fila omitida: {exc}")
                continue
            if campos is not None:
                candidatas.append(campos)

        if not candidatas:
            return 0

        identificaciones = list({c["identificacion"] for c in candidatas})
        existentes = await siniestro_previsional_repository.get_by_identificaciones(
            db, identificaciones
        )
        claves_existentes = {(s.identificacion, s.numero_siniestro) for s in existentes}

        a_insertar: list[dict[str, Any]] = []
        vistas: set[tuple[str, str]] = set()
        for c in candidatas:
            clave = (c["identificacion"], c["numero_siniestro"])
            if clave in claves_existentes or clave in vistas:
                continue
            vistas.add(clave)
            a_insertar.append(c)

        creados = await siniestro_previsional_repository.bulk_create_flushed(db, a_insertar)
        await db.commit()
        return len(creados)

    # -- SOLICITUDES ------------------------------------------------------

    @staticmethod
    def _parsear_fila_solicitud(fila: tuple, num_fila: int) -> dict[str, Any] | None:
        """
        Mapeo SOLICITUDES: D=identificacion, B=radicado (interpretado como
        "No. Solicitud Auditoria" — es la columna más cercana a "número de
        radicado de la solicitud en ARPIS" que trae la hoja; ver reporte de
        Task 3.2), N=dia_181, R=observacion, AF=fecha_crie.

        `SolicitudPrevisional.tipo_ingreso`, `.fecha_inicial` y
        `.fecha_final` quedan SIN poblar: la hoja SOLICITUDES no trae
        ninguna columna con un mapeo obvio para esos tres campos (las
        fechas inicial/final de la incapacidad viven en la hoja "LISTADO
        ITE DIA", no en SOLICITUDES) — GAP documentado en el reporte, no
        se adivina el mapeo.

        Retorna `None` (fila omitida sin contar como error) si falta la
        identificación. Propaga `CeldaInvalidaError` si una fecha no se
        puede parsear.
        """
        identificacion = texto(_get(fila, 3))  # D
        radicado = texto(_get(fila, 1))  # B
        dia_181 = parse_fecha(_get(fila, 13))  # N
        observacion = texto(_get(fila, 17))  # R
        fecha_crie = parse_fecha(_get(fila, 31))  # AF

        if es_vacio(identificacion):
            logger.warning(
                f"SOLICITUDES fila {num_fila}: identificación vacía — fila "
                "omitida (requerida por SolicitudPrevisional.identificacion)."
            )
            return None

        return {
            "identificacion": identificacion,
            "radicado": radicado,
            "dia_181": dia_181,
            "fecha_crie": fecha_crie,
            "observacion": observacion,
        }

    async def importar_solicitudes(self, db: AsyncSession, file_bytes: bytes) -> int:
        ws = _cargar_hoja(file_bytes, _HOJA_SOLICITUDES)
        validar_encabezados(ws, _ENCABEZADOS_SOLICITUDES, fila=_FILA_ENCABEZADO)

        candidatas: list[dict[str, Any]] = []
        for i, fila in enumerate(
            ws.iter_rows(min_row=_FILA_ENCABEZADO + 1, values_only=True),
            start=_FILA_ENCABEZADO + 1,
        ):
            try:
                campos = self._parsear_fila_solicitud(fila, i)
            except CeldaInvalidaError as exc:
                logger.warning(f"SOLICITUDES fila {i}: error de parseo, fila omitida: {exc}")
                continue
            if campos is not None:
                candidatas.append(campos)

        if not candidatas:
            return 0

        identificaciones = list({c["identificacion"] for c in candidatas})
        existentes = await solicitud_previsional_repository.get_by_identificaciones(
            db, identificaciones
        )
        # Solo se puede deduplicar de forma confiable cuando hay `radicado`
        # — sin él no hay llave natural (ver docstring del módulo).
        claves_existentes = {
            (s.identificacion, s.radicado) for s in existentes if s.radicado is not None
        }

        a_insertar: list[dict[str, Any]] = []
        vistas: set[tuple[str, str]] = set()
        for c in candidatas:
            if c["radicado"] is None:
                a_insertar.append(c)
                continue
            clave = (c["identificacion"], c["radicado"])
            if clave in claves_existentes or clave in vistas:
                continue
            vistas.add(clave)
            a_insertar.append(c)

        creadas = await solicitud_previsional_repository.bulk_create_flushed(db, a_insertar)
        await db.commit()
        return len(creadas)

    # -- LISTADO ITE DIA --------------------------------------------------

    @staticmethod
    def _parsear_fila_ite(fila: tuple, num_fila: int) -> dict[str, Any] | None:
        """
        Mapeo "LISTADO ITE DIA": F=identificacion, I=fecha_inicial (junto
        con F forman la llave anti-doble-pago de la regla AT),
        J=fecha_final. La columna A ("FI") es una llave concatenada
        precalculada en la fuente — se IGNORA por completo, siguiendo el
        principio ya establecido en todo este módulo de nunca confiar en
        claves-string concatenadas: la llave tipada
        `(identificacion, fecha_inicial)` se construye acá desde F/I.

        `IteHistorico.valor_pagado` y `.fecha_pago` quedan SIN poblar: la
        hoja trae varias columnas de valor (AFP, Alfa, por grupo salarial)
        y de fecha (radicación, radicación ALFA, auditoría) pero ninguna
        es inequívocamente "el valor pagado" / "la fecha de pago" — GAP
        documentado en el reporte, no se adivina cuál usar.

        Retorna `None` (fila omitida sin contar como error) si falta
        identificación o fecha inicial (ambas requeridas por el modelo —
        `fecha_inicial` es NOT NULL). Propaga `CeldaInvalidaError` si una
        fecha no se puede parsear.
        """
        identificacion = texto(_get(fila, 5))  # F
        fecha_inicial = parse_fecha(_get(fila, 8))  # I
        fecha_final = parse_fecha(_get(fila, 9))  # J

        if es_vacio(identificacion) or fecha_inicial is None:
            logger.warning(
                f"LISTADO ITE DIA fila {num_fila}: identificación o fecha "
                "inicial vacía — fila omitida (requeridas por "
                "IteHistorico.identificacion/fecha_inicial)."
            )
            return None

        return {
            "identificacion": identificacion,
            "fecha_inicial": fecha_inicial,
            "fecha_final": fecha_final,
        }

    async def importar_ite_historico(self, db: AsyncSession, file_bytes: bytes) -> int:
        ws = _cargar_hoja(file_bytes, _HOJA_ITE)
        validar_encabezados(ws, _ENCABEZADOS_ITE, fila=_FILA_ENCABEZADO)

        candidatas: list[dict[str, Any]] = []
        for i, fila in enumerate(
            ws.iter_rows(min_row=_FILA_ENCABEZADO + 1, values_only=True),
            start=_FILA_ENCABEZADO + 1,
        ):
            try:
                campos = self._parsear_fila_ite(fila, i)
            except CeldaInvalidaError as exc:
                logger.warning(f"LISTADO ITE DIA fila {i}: error de parseo, fila omitida: {exc}")
                continue
            if campos is not None:
                candidatas.append(campos)

        if not candidatas:
            return 0

        claves_candidatas = list(
            {(c["identificacion"], c["fecha_inicial"]) for c in candidatas}
        )
        claves_existentes = await ite_historico_repository.get_existentes(
            db, claves_candidatas
        )

        a_insertar: list[dict[str, Any]] = []
        vistas: set[tuple[str, Any]] = set()
        for c in candidatas:
            clave = (c["identificacion"], c["fecha_inicial"])
            if clave in claves_existentes or clave in vistas:
                continue
            vistas.add(clave)
            a_insertar.append(c)

        creados = await ite_historico_repository.bulk_create_flushed(db, a_insertar)
        await db.commit()
        return len(creados)


# Instancia global del adaptador (mismo patrón que
# `app/services/integracion/servialfa_client.py` / `sicat_client.py` y el
# resto de servicios previsionales — singleton a nivel de módulo).
excel_referencia_adapter = ExcelReferenciaAdapter()
